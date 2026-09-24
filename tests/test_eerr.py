"""
Verificación de "EERR Dinámico" (24-09-2026) — Estado de Resultados por
conceptos a partir del Estado de Resultados Comparativo que deja el puente
Nubox -> repo en `nubox_importado/estado_resultado_comparativo/`. Mismo
estilo sin-pytest que `tests/test_analisis.py` (reusa el bootstrap de
`tests/run_verification.py` — `flask_app`).

Uso:
    python tests/test_eerr.py
"""

import csv
import io
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests.run_verification import flask_app, seed_data  # noqa: E402
from app.eerr import calculo  # noqa: E402
from app.data import nubox_importado_repo  # noqa: E402

PASSED, FAILED = [], []


def check(label, condition, extra=""):
    if condition:
        PASSED.append(label)
        print(f"  OK  {label}")
    else:
        FAILED.append(label)
        print(f" FAIL {label} {extra}")


ENCABEZADOS = ["idcuenta", "tipocuentaid", "tipocuenta", "cuentamayor", "subcuenta", "cuentaanalisis",
               "periodo", "saldoanterior", "saldo", "generado_en"]


def _fila(tipo, mayor, sub, mes, saldo, analisis=""):
    return ["1", str(tipo), "Total Ganancias" if tipo == 3 else "Total Pérdidas", mayor, sub, analisis,
            f"2026-{mes:02d}-01T00:00:00.000Z", "0", str(saldo), "2026-09-24T03:00:00.000Z"]


def _csv_ejemplo(path):
    filas = [
        _fila(3, "5101 INGRESOS", "", 1, 1500),  # subtotal de cuenta mayor: se ignora
        _fila(3, "5101 INGRESOS", "5101-01 VENTAS", 1, 1000),
        _fila(3, "5101 INGRESOS", "5101-01 VENTAS", 2, 2000),
        _fila(3, "5101 INGRESOS", "5101-02 OTRAS  VENTAS", 1, 500),
        _fila(3, "5101 INGRESOS", "5101-02 OTRAS  VENTAS", 1, 999, analisis="X-1 DETALLE"),  # detalle: se ignora
        _fila(4, "4101 COSTOS", "4101-01 COSTO DE VENTA", 1, 400),
        _fila(4, "4101 COSTOS", "4101-01 COSTO DE VENTA", 2, 800),
        _fila(4, "4201 GASTOS", "4201-01 REMUNERACIONES", 2, 300),
        _fila(4, "4201 GASTOS", "4201-09 SIN MOVIMIENTO", 1, 0),
        _fila(4, "4701 IMPUESTOS", "4701-01 IMPTO PRIMERA CATEGORIA", 2, 100),
    ]
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(ENCABEZADOS)
        w.writerows(filas)


def test_lectura_y_calculo(tmp_dir):
    ruta = tmp_dir / "x_2026.csv"
    _csv_ejemplo(ruta)
    datos = calculo.leer_eerr_csv(ruta)
    cuentas = {c["codigo"]: c for c in datos["cuentas"]}
    check("lee solo subcuentas con movimiento", sorted(cuentas) == ["4101-01", "4201-01", "4701-01", "5101-01", "5101-02"], str(sorted(cuentas)))
    check("ignora el nivel de cuenta de análisis", cuentas["5101-02"]["meses"][0] == 500)
    check("nombre sin espacios dobles", cuentas["5101-02"]["nombre"] == "OTRAS VENTAS")
    check("meses con datos", datos["meses_con_datos"] == [1, 2])
    check("tipo ganancia/pérdida", cuentas["5101-01"]["tipo"] == 3 and cuentas["4101-01"]["tipo"] == 4)

    conceptos = calculo.conceptos_por_defecto()
    asig = {"5101-01": "ing_exp", "5101-02": "ing_exp", "4101-01": "costos_exp", "4201-01": "remu"}
    inf = calculo.calcular(datos["cuentas"], conceptos, asig, 1, 2)
    filas = {f.get("id") or f["nombre"]: f for f in inf["filas"] if f["tipo"] != "cta"}
    check("columnas del rango", inf["columnas"] == ["Enero", "Febrero"])
    check("Total Ganancias por mes", filas["tg"]["valores"] == [1500, 2000], str(filas["tg"]["valores"]))
    check("Margen Operacional = TG − TC", filas["mo"]["valores"] == [1100, 1200])
    check("Margen Bruto = MO − TOC", filas["mb"]["valores"] == [1100, 900])
    check("Resultado = MB − TE", filas["res"]["valores"] == [1100, 900])
    check("% de concepto sobre Total Ganancias", round(filas["costos_exp"]["pcts"][0], 4) == round(400 / 1500 * 100, 4))
    ctas = [f for f in inf["filas"] if f["tipo"] == "cta" and f["padre"] == "ing_exp"]
    check("% de cuenta sobre su concepto", round(ctas[0]["pcts"][0], 4) == round(1000 / 1500 * 100, 4))
    check("cuenta sin asignar se informa", inf["sin_asignar"] == ["4701-01"])
    check("fila de efecto sin asignar", any(f["tipo"] == "sinasig" and f["valores"] == [0, -100] for f in inf["filas"]))
    check("diferencia con Nubox = cuenta sin asignar", inf["cuadre"]["diferencia"] == 100 and not inf["cuadre"]["cuadra"])

    asig["4701-01"] = "renta"
    inf = calculo.calcular(datos["cuentas"], conceptos, asig, 1, 2)
    check("con todo asignado cuadra con Nubox", inf["cuadre"]["cuadra"] and inf["cuadre"]["resultado_informe"] == 1900)

    # Ganancia puesta en una sección de pérdidas: resta, y el resultado sigue cuadrando.
    asig2 = dict(asig, **{"5101-02": "ocp"})
    inf2 = calculo.calcular(datos["cuentas"], conceptos, asig2, 1, 2)
    f2 = {f.get("id"): f for f in inf2["filas"] if f.get("id")}
    check("cuenta de sección contraria resta", f2["ocp"]["valores"] == [-500, 0])
    check("y el resultado sigue cuadrando", inf2["cuadre"]["cuadra"])

    # "No incluir" oculta la fila pero no cambia los totales.
    sin_remu = [dict(c, inc=(c["id"] != "remu")) for c in conceptos]
    inf3 = calculo.calcular(datos["cuentas"], sin_remu, asig, 1, 2)
    ids3 = [f.get("id") for f in inf3["filas"]]
    f3 = {f.get("id"): f for f in inf3["filas"] if f.get("id")}
    check("'No incluir' oculta la fila", "remu" not in ids3 and not any(f.get("padre") == "remu" for f in inf3["filas"]))
    check("'No incluir' mantiene el total", f3["toc"]["valores"] == [0, 300])

    inf4 = calculo.calcular(datos["cuentas"], conceptos, asig, 2, 2)
    check("rango de un mes", inf4["columnas"] == ["Febrero"] and [f for f in inf4["filas"] if f.get("id") == "tg"][0]["valores"] == [2000])


def test_normalizacion():
    enviados = [
        {"id": "tg", "tipo": "t", "inc": False},
        {"id": "mo", "tipo": "f", "a": "res", "b": "tg"},  # intento de alterar la fórmula
        {"id": "c1", "n": "  publicidad   y mkt ", "tipo": "g", "sec": "toc"},
        {"id": "c2", "n": "mal", "tipo": "g", "sec": "no-existe"},
        {"id": "../x", "n": "raro", "tipo": "g", "sec": "tg"},
    ]
    res = calculo.normalizar_conceptos(enviados)
    ids = [c["id"] for c in res]
    check("mantiene todos los conceptos base", all(c["id"] in ids for c in calculo.CONCEPTOS_BASE))
    check("respeta 'No incluir' de un total", [c for c in res if c["id"] == "tg"][0]["inc"] is False)
    check("no deja alterar fórmulas", [c for c in res if c["id"] == "mo"][0]["a"] == "tg")
    check("concepto propio en mayúsculas y antes de su total", ids.index("c1") == ids.index("toc") - 1
          and [c for c in res if c["id"] == "c1"][0]["n"] == "PUBLICIDAD Y MKT")
    check("descarta secciones o ids inválidos", "c2" not in ids and "../x" not in ids)
    asig = calculo.normalizar_asignaciones({"5101-01": "ing_exp", "5101-02": "tg", "../": "remu", "4101-01": "c1"}, res)
    check("asignaciones solo a conceptos que agrupan", asig == {"5101-01": "ing_exp", "4101-01": "c1"}, str(asig))


def test_datos_reales_del_puente():
    empresas = nubox_importado_repo.listar_empresas_nubox()
    check("empresas del conector (puente)", {e["alias"] for e in empresas} >= {"626"}, str(empresas))
    check("años disponibles 626", 2026 in nubox_importado_repo.eerr_anios_disponibles("626"))
    check("alias inválido no arma ruta", nubox_importado_repo.ruta_eerr("../626", 2026) is None)
    datos = calculo.leer_eerr_csv(nubox_importado_repo.ruta_eerr("626", 2026))
    ventas = [c for c in datos["cuentas"] if c["codigo"] == "5101-58"]
    check("626: 5101-58 VENTA DE PRODUCTO ENTERO enero", ventas and ventas[0]["meses"][0] == 38441484)
    inf = calculo.calcular(datos["cuentas"], calculo.conceptos_por_defecto(), {}, 1, 8)
    check("626: sin asignar, diferencia = resultado de Nubox", abs(inf["cuadre"]["diferencia"] + inf["cuadre"]["resultado_nubox"]) < 1)


def _login_admin():
    client = flask_app.test_client()
    client.post("/login", data={"usuario": "", "clave": "test_local_only_1234"}, follow_redirects=True)
    return client


def test_rutas_flujo_completo():
    client = _login_admin()
    r = client.get("/eerr/")
    html = r.get_data(as_text=True)
    check("GET /eerr/ 200", r.status_code == 200)
    check("página lista empresas de Nubox", "SOCIEDAD AGRICOLA Y COMERCIAL ACEVEDO" in html)
    check("menú muestra EERR Dinámico", "EERR Dinámico" in html)

    r = client.get("/eerr/datos?sistema=nubox&codigo=626&anio=2026")
    j = r.get_json()
    check("GET /eerr/datos 200", r.status_code == 200, str(j)[:200])
    check("datos: meses enero-agosto", j["meses_con_datos"] == list(range(1, 9)))
    check("datos: plantilla de conceptos", [c["id"] for c in j["conceptos"]][:4] == ["ing_exp", "otros_ing", "ing_reaj", "tg"])
    check("datos: sin configuración guardada", j["config_guardada"] is False)

    r = client.get("/eerr/datos?sistema=nubox&codigo=../626&anio=2026")
    check("datos: empresa inexistente 404", r.status_code == 404)
    r = client.get("/eerr/datos?sistema=softland&codigo=626&anio=2026")
    check("datos: softland sin empresa 404", r.status_code == 404)

    conceptos = j["conceptos"] + []
    conceptos.insert(3, {"id": "c99", "n": "Ventas especiales", "tipo": "g", "sec": "tg", "inc": True, "custom": True})
    asig = {"5101-58": "ing_exp", "5101-59": "c99"}
    cuerpo = {"sistema": "nubox", "codigo": "626", "anio": 2026, "desde": 1, "hasta": 8,
              "conceptos": conceptos, "asignaciones": asig, "guardar": True}
    r = client.post("/eerr/informe", data=json.dumps(cuerpo), content_type="application/json")
    j2 = r.get_json()
    check("POST /eerr/informe 200", r.status_code == 200, str(j2)[:200])
    check("informe guardado", j2["guardado"] is True and j2["aviso"] is None)
    fila_c99 = [f for f in j2["informe"]["filas"] if f.get("id") == "c99"]
    check("concepto propio en el informe", fila_c99 and fila_c99[0]["valores"][0] == 50137920)

    r = client.get("/eerr/datos?sistema=nubox&codigo=626&anio=2026")
    j3 = r.get_json()
    check("configuración queda guardada", j3["config_guardada"] and j3["asignaciones"] == asig)
    check("concepto propio queda guardado", any(c["id"] == "c99" and c["n"] == "VENTAS ESPECIALES" for c in j3["conceptos"]))

    form = {"sistema": "nubox", "codigo": "626", "anio": "2026", "desde": "1", "hasta": "8",
            "conceptos": json.dumps(j3["conceptos"]), "asignaciones": json.dumps(asig), "con_cuentas": "1"}
    r = client.post("/eerr/exportar", data=dict(form, formato="xlsx"))
    check("exporta Excel", r.status_code == 200 and "spreadsheetml" in (r.headers.get("Content-Type") or ""))
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb.active
    textos = [ws.cell(row=i, column=1).value for i in range(1, ws.max_row + 1)]
    check("Excel: encabezado y filas", ws["A2"].value.startswith("SOCIEDAD AGRICOLA") and "Total Ganancias" in textos
          and "RESULTADO DEL EJERCICIO" in textos, str(textos[:8]))
    check("Excel: columnas Enero 2026 y %", ws.cell(row=5, column=2).value == "ENERO 2026" and ws.cell(row=5, column=3).value == "%")
    fila_venta = textos.index("    5101-58 VENTA DE PRODUCTO ENTERO") + 1
    check("Excel: monto de la cuenta", ws.cell(row=fila_venta, column=2).value == 38441484)

    r = client.post("/eerr/exportar", data=dict(form, formato="pdf", con_cuentas="0"))
    check("exporta PDF", r.status_code == 200 and r.data[:4] == b"%PDF")
    check("nombre de archivo del PDF", "EERR_626_202601-08.pdf" in (r.headers.get("Content-Disposition") or ""))


def test_admin_sistemas_contables():
    client = _login_admin()
    r = client.get("/admin/sistemas_contables")
    html = r.get_data(as_text=True)
    check("GET admin sistemas contables 200", r.status_code == 200)
    check("admin lista empresas del conector Nubox", "NEXT PRO SPA" in html)
    r = client.post("/admin/sistemas_contables/agregar",
                    data={"codigo": "701", "razon_social": "EMPRESA  DEMO SPA", "rut": "76.000.000-0", "sistema": "softland"},
                    follow_redirects=True)
    check("agrega empresa Softland", "EMPRESA DEMO SPA agregada a Softland" in r.get_data(as_text=True))
    r = client.post("/admin/sistemas_contables/agregar",
                    data={"codigo": "701", "razon_social": "OTRA", "sistema": "softland"}, follow_redirects=True)
    check("no duplica código en el mismo sistema", "Ya existe la empresa 701" in r.get_data(as_text=True))
    r = client.post("/admin/sistemas_contables/agregar",
                    data={"codigo": "702", "razon_social": "X", "sistema": "nubox"}, follow_redirects=True)
    check("no permite agregar a mano en Nubox", "Elige Softland o Defontana" in r.get_data(as_text=True))
    r = client.post("/admin/sistemas_contables/agregar",
                    data={"codigo": "7 01/..", "razon_social": "X", "sistema": "defontana"}, follow_redirects=True)
    check("valida el código", "El código solo puede tener" in r.get_data(as_text=True))

    r = client.get("/eerr/")
    check("empresa Softland aparece en EERR", "EMPRESA DEMO SPA" in r.get_data(as_text=True))
    r = client.get("/eerr/datos?sistema=softland&codigo=701&anio=2026")
    check("Softland: importación por definir", r.status_code == 404 and "aún no está disponible" in r.get_json()["error"])

    from app.data import eerr_repo
    emp = [e for e in eerr_repo.listar_empresas_manuales() if e["codigo"] == "701"][0]
    client.post("/admin/sistemas_contables/eliminar", data={"id": emp["id"]}, follow_redirects=True)
    check("quita empresa", not any(e["codigo"] == "701" for e in eerr_repo.listar_empresas_manuales()))


def test_trabajador_no_puede_ver_eerr():
    client = flask_app.test_client()
    client.post("/login", data={"usuario": "testuser", "clave": "trabajador123"}, follow_redirects=True)
    r = client.get("/eerr/")
    check("trabajador NO ve EERR Dinámico (403)", r.status_code == 403)
    r = client.get("/admin/sistemas_contables")
    check("trabajador NO ve Sistemas contables", r.status_code in (302, 403))


def main():
    import tempfile

    seed_data()

    with tempfile.TemporaryDirectory() as tmp:
        test_lectura_y_calculo(Path(tmp))
    test_normalizacion()
    test_datos_reales_del_puente()
    test_rutas_flujo_completo()
    test_admin_sistemas_contables()
    test_trabajador_no_puede_ver_eerr()

    print(f"\n{len(PASSED)} OK, {len(FAILED)} FAIL")
    if FAILED:
        print("Fallaron:", FAILED)
        sys.exit(1)


if __name__ == "__main__":
    main()
