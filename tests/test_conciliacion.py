"""
Verificación de "Conciliación" — conciliación asistida contra documentos
auxiliares vía el modal "Crear comprobante" (14-09-2026), mismo estilo
sin-pytest que `tests/test_sii.py`/`tests/test_depreciacion.py` (reusa el
bootstrap de `tests/run_verification.py` — `flask_app` — importándolo
como módulo).

Todo el matching automático (monto Y rut) y la interacción del modal
(agregar/quitar líneas, adjuntar documentos, buscador combinado) viven en
`app/static/js/conciliacion.js` (cliente puro) — no se pueden probar acá,
donde solo corre Python. Estas pruebas cubren todo lo que SÍ es del
servidor: parseo de los 3 Excel auxiliares, la ruta de carga AJAX, la
reconstrucción del pool de auxiliares Y de las líneas ya armadas tras un
error de validación, y el armado del comprobante final — cada línea sin
documentos aporta una fila, cada línea CON documentos se expande en una
fila por documento (mismo código de cuenta, cada una con su propio bloque
de Tipo Auxiliar "A"/"H") — balance Debe==Haber en cualquier combinación
de líneas/documentos.

Uso:
    python tests/test_conciliacion.py
"""

import io
import sys
from datetime import datetime
from pathlib import Path

import xlrd
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests.run_verification import flask_app  # noqa: E402
from app.conciliacion import export_writer  # noqa: E402
from app.conciliacion.documentos import AUXILIAR_MODULOS  # noqa: E402
from app.conciliacion.plan_cuentas import CUENTAS_POR_CODIGO  # noqa: E402

PASSED, FAILED = [], []


def check(label, condition, extra=""):
    if condition:
        PASSED.append(label)
        print(f"  OK  {label}")
    else:
        FAILED.append(label)
        print(f" FAIL {label} {extra}")


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _xlsx_cartola(filas):
    """`filas`: lista de (fecha_dt, detalle, cargo, abono). Mismo layout
    que arma `app/parsers/output_writer.py` (hoja "Banco", columnas
    B:E)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Banco"
    ws["B1"] = "FECHA DIA/MES"
    ws["C1"] = "DETALLE DE TRANSACCION"
    ws["D1"] = "MONTO CHEQUES O CARGOS"
    ws["E1"] = "MONTO DEPOSITOS O ABONOS"
    for i, (fecha, detalle, cargo, abono) in enumerate(filas, start=2):
        ws.cell(row=i, column=2, value=fecha)
        ws.cell(row=i, column=3, value=detalle)
        ws.cell(row=i, column=4, value=cargo or None)
        ws.cell(row=i, column=5, value=abono or None)
    return _bytes(wb)


def _xlsx_clientes_proveedores(docs, col_monto):
    """`docs`: lista de dicts {rut_numero, rut_dv, nombre, fecha (datetime),
    tipo_documento, numero_documento, monto}. `col_monto`: 16 (P, Clientes)
    o 17 (Q, Proveedores) — mismo layout que `caja_parsers.parsear_
    clientes`/`parsear_proveedores` (encabezado fila 2, datos desde fila 3)."""
    wb = Workbook()
    ws = wb.active
    for i, d in enumerate(docs, start=3):
        ws.cell(row=i, column=3, value=d["rut_numero"])
        ws.cell(row=i, column=4, value=d["rut_dv"])
        ws.cell(row=i, column=5, value=d["nombre"])
        ws.cell(row=i, column=6, value=d["fecha"])
        ws.cell(row=i, column=11, value=d["tipo_documento"])
        ws.cell(row=i, column=12, value=d["numero_documento"])
        ws.cell(row=i, column=col_monto, value=d["monto"])
    return _bytes(wb)


def _xlsx_honorarios(docs):
    """`docs`: lista de dicts {rut, nombre, fecha_txt (DD-MM-AAAA), boleta,
    monto} — mismo layout que `caja_parsers.parsear_honorarios` (encabezado
    fila 3, datos desde fila 4)."""
    wb = Workbook()
    ws = wb.active
    for i, d in enumerate(docs, start=4):
        ws.cell(row=i, column=1, value=d["rut"])
        ws.cell(row=i, column=2, value=d["nombre"])
        ws.cell(row=i, column=3, value=d["fecha_txt"])
        ws.cell(row=i, column=6, value=d["boleta"])
        ws.cell(row=i, column=10, value=d["monto"])
    return _bytes(wb)


def test_procesar_fecha_guardada_como_texto():
    """Regresión (14-09-2026, reportado por el usuario con un caso real):
    si "Subir Cartolas" no pudo reconocer el formato original de una
    fecha y la dejó como TEXTO plano en la celda (`app/parsers/output_
    writer.py:_parse_fecha`) en vez de una fecha real de Excel, se veía
    bien en pantalla ("12-08-2026") pero `fecha_iso` quedaba vacío y
    bloqueaba la descarga con "no trae una fecha reconocible" pese a que
    el texto SÍ es una fecha válida — `_leer_excel_convertido` debe
    reintentar interpretarlo como texto antes de rendirse."""
    client = flask_app.test_client()
    client.post("/login", data={"usuario": "", "clave": "test_local_only_1234"}, follow_redirects=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Banco"
    ws.cell(row=2, column=2, value="12-08-2026")  # fecha como TEXTO, no datetime
    ws.cell(row=2, column=3, value="OF VIRT U")
    ws.cell(row=2, column=5, value=912968)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    r = client.post(
        "/conciliacion/procesar",
        data={"archivo": (buf, "cartola.xlsx")},
        content_type="multipart/form-data",
    )
    body = r.get_data(as_text=True)
    check("procesar con fecha en texto -> 200", r.status_code == 200)
    check("la fecha se sigue mostrando igual (12-08-2026)", "12-08-2026" in body)
    check("fecha_iso_0 quedó poblado (2026-08-12), no vacío", 'name="fecha_iso_0" value="2026-08-12"' in body)


def test_auxiliar_modulos_config():
    check("Clientes: cuenta fija 1104-01, bloque A, dirección abono",
          AUXILIAR_MODULOS["clientes"]["cuenta_codigo"] == "1104-01"
          and AUXILIAR_MODULOS["clientes"]["tipo_auxiliar"] == "A"
          and AUXILIAR_MODULOS["clientes"]["direccion"] == "abono")
    check("Proveedores: cuenta fija 2105-01, bloque A, dirección cargo",
          AUXILIAR_MODULOS["proveedores"]["cuenta_codigo"] == "2105-01"
          and AUXILIAR_MODULOS["proveedores"]["tipo_auxiliar"] == "A"
          and AUXILIAR_MODULOS["proveedores"]["direccion"] == "cargo")
    check("Honorarios: cuenta fija 2105-04, bloque H, dirección cargo",
          AUXILIAR_MODULOS["honorarios"]["cuenta_codigo"] == "2105-04"
          and AUXILIAR_MODULOS["honorarios"]["tipo_auxiliar"] == "H"
          and AUXILIAR_MODULOS["honorarios"]["direccion"] == "cargo")
    for modulo, info in AUXILIAR_MODULOS.items():
        check(f"{modulo}: la cuenta fija existe en el plan de cuentas", info["cuenta_codigo"] in CUENTAS_POR_CODIGO)


CUENTA_BANCO = {"codigo": "1101-29", "descripcion": "BANCO BCI", "es_banco": True, "requiere_centro_costo": False}
CUENTA_CLIENTES = {"codigo": "1104-01", "descripcion": "DEUDORES CLIENTES", "es_banco": False, "requiere_centro_costo": False}
CUENTA_GASTO = {"codigo": "5501-05", "descripcion": "GASTOS VARIOS", "es_banco": False, "requiere_centro_costo": False}
CUENTA_OTRO_GASTO = {"codigo": "4201-11", "descripcion": "MULTAS", "es_banco": False, "requiere_centro_costo": False}


def test_construir_filas_comprobantes_linea_plana():
    """Una línea sin documentos aporta UNA fila con el monto digitado a
    mano — comportamiento base, sin auxiliar."""
    movimientos = [{
        "fecha": datetime(2026, 8, 21), "detalle": "OF CENTRA", "cargo": 50_000.0, "abono": 0.0,
        "lineas": [{"cuenta": CUENTA_GASTO, "monto": 50_000.0, "documentos": []}],
    }]
    filas = export_writer.construir_filas_comprobantes(movimientos, CUENTA_BANCO)
    check("2 filas (banco + 1 línea)", len(filas) == 2)
    fila_banco = next(f for f in filas if f[4] == "1101-29")
    fila_concepto = next(f for f in filas if f[4] == "5501-05")
    check("banco con Tipo Auxiliar 'B'", fila_banco[10] == "B")
    check("línea plana sin Tipo Auxiliar", not fila_concepto[10])
    check("balanceado", sum(f[8] for f in filas if f[8]) == sum(f[9] for f in filas if f[9]) == 50_000.0)
    check("primera fila (Número=0/Tipo/Fecha/Glosa) es la línea de detalle en un cargo", fila_concepto[0] == 0 and fila_concepto[1] == "E")
    check("la fila del banco no repite Número/Tipo/Fecha/Glosa", fila_banco[0] == "" and fila_banco[1] == "")


def test_construir_filas_comprobantes_linea_con_varios_documentos():
    """Una línea CON documentos (14-09-2026, modal "Crear comprobante") se
    expande en una fila POR documento — mismo código de cuenta, cada una
    con su propio bloque de Tipo Auxiliar "A", y el comprobante sigue
    balanceado contra el total del movimiento."""
    doc_a = {"tipo": "A", "rut": "1-9", "nombre": "Doc A", "tipo_documento_codigo": 33, "numero_documento": "1", "fecha": datetime(2026, 8, 1), "monto": 60_000.0}
    doc_b = {"tipo": "A", "rut": "2-7", "nombre": "Doc B", "tipo_documento_codigo": 33, "numero_documento": "2", "fecha": datetime(2026, 8, 2), "monto": 43_733.0}
    movimientos = [{
        "fecha": datetime(2026, 8, 18), "detalle": "TRANSFERENCIA VARIOS CLIENTES", "cargo": 0.0, "abono": 103_733.0,
        "lineas": [{"cuenta": CUENTA_CLIENTES, "monto": 0, "documentos": [doc_a, doc_b]}],
    }]
    filas = export_writer.construir_filas_comprobantes(movimientos, CUENTA_BANCO)
    check("3 filas (banco + 2 documentos)", len(filas) == 3)
    check("es un abono -> la primera fila es el banco", filas[0][4] == "1101-29" and filas[0][0] == 0 and filas[0][1] == "I")

    filas_doc = [f for f in filas if f[4] == "1104-01"]
    check("2 filas de la cuenta Clientes (una por documento)", len(filas_doc) == 2)
    check("cada fila lleva el rut de SU documento", {f[11] for f in filas_doc} == {"1-9", "2-7"})
    check("cada fila lleva el monto de SU documento (no el total combinado)", {f[9] for f in filas_doc} == {60_000.0, 43_733.0})
    check("balanceado", sum(f[8] for f in filas if f[8]) == sum(f[9] for f in filas if f[9]) == 103_733.0)


def test_construir_filas_comprobantes_varias_lineas():
    """"+ Agregar cuenta" del modal: un movimiento partido en más de una
    línea (ninguna con documentos) — cada línea aporta su propia fila."""
    movimientos = [{
        "fecha": datetime(2026, 8, 20), "detalle": "PAGO MIXTO", "cargo": 161_721.0, "abono": 0.0,
        "lineas": [
            {"cuenta": CUENTA_GASTO, "monto": 100_000.0, "documentos": []},
            {"cuenta": CUENTA_OTRO_GASTO, "monto": 61_721.0, "documentos": []},
        ],
    }]
    filas = export_writer.construir_filas_comprobantes(movimientos, CUENTA_BANCO)
    check("3 filas (banco + 2 líneas)", len(filas) == 3)
    check("balanceado", sum(f[8] for f in filas if f[8]) == sum(f[9] for f in filas if f[9]) == 161_721.0)
    check("la primera línea de detalle es la primera fila (cargo)", filas[0][4] == "5501-05" and filas[0][0] == 0)


def test_cargar_auxiliar_rutas():
    client = flask_app.test_client()
    r = client.post("/login", data={"usuario": "", "clave": "test_local_only_1234"}, follow_redirects=True)
    check("login admin ok (fixture)", r.status_code == 200)

    r = client.post("/conciliacion/cargar-auxiliar/modulo-inexistente", data={}, content_type="multipart/form-data")
    check("módulo desconocido -> 404", r.status_code == 404)

    r = client.post("/conciliacion/cargar-auxiliar/clientes", data={}, content_type="multipart/form-data")
    check("sin archivo -> 400", r.status_code == 400)

    xlsx_clientes = _xlsx_clientes_proveedores(
        [{"rut_numero": "76123456", "rut_dv": "7", "nombre": "PAC HDI SEGUROS SA", "fecha": datetime(2026, 8, 1),
          "tipo_documento": "FAC-EL", "numero_documento": "456", "monto": 103_733}],
        col_monto=16,
    )
    r = client.post(
        "/conciliacion/cargar-auxiliar/clientes",
        data={"archivo": (io.BytesIO(xlsx_clientes), "clientes.xlsx")},
        content_type="multipart/form-data",
    )
    body = r.get_data(as_text=True)
    check("cargar Clientes -> 200", r.status_code == 200)
    check("el fragmento trae el nombre del documento", "PAC HDI SEGUROS SA" in body)
    check("el fragmento trae el RUT armado (76123456-7)", "76123456-7" in body)
    check("el fragmento trae el monto formateado", "103.733" in body or "103733" in body)
    check("el fragmento marca el documento como 'Disponible'", "Disponible" in body)
    check("el fragmento trae el campo oculto aux_clientes_total = 1", 'name="aux_clientes_total" id="aux-total-clientes" value="1"' in body)

    xlsx_vacio = _xlsx_clientes_proveedores([], col_monto=16)
    r = client.post(
        "/conciliacion/cargar-auxiliar/clientes",
        data={"archivo": (io.BytesIO(xlsx_vacio), "vacio.xlsx")},
        content_type="multipart/form-data",
    )
    check("Excel sin documentos -> 400", r.status_code == 400)

    xlsx_honorarios = _xlsx_honorarios([
        {"rut": "16876802-8", "nombre": "Juan Pérez", "fecha_txt": "10-08-2026", "boleta": "BOL-HE 7", "monto": 80_000},
    ])
    r = client.post(
        "/conciliacion/cargar-auxiliar/honorarios",
        data={"archivo": (io.BytesIO(xlsx_honorarios), "honorarios.xlsx")},
        content_type="multipart/form-data",
    )
    body_hon = r.get_data(as_text=True)
    check("cargar Honorarios -> 200", r.status_code == 200)
    check("Honorarios: separa boleta en tipo BOL-HE + número 7", 'value="BOL-HE"' in body_hon and 'value="7"' in body_hon)


def test_descargar_con_lineas_multiples():
    """Flujo completo de `/conciliacion/descargar` con 3 movimientos —
    exactamente lo que dejaría armado el modal "Crear comprobante" antes
    de enviar el formulario:

    - Fila 0 (abono): 1 línea (Clientes) con 1 documento adjunto.
    - Fila 1 (cargo): 2 líneas — una (Proveedores) con 1 documento que NO
      cubre el monto completo del movimiento, y otra plana (cuenta
      suelta) con el resto ("+ Agregar cuenta").
    - Fila 2 (cargo): 1 línea plana, sin documentos.

    Verifica que el .xls final balancea y que cada fila trae el bloque de
    Tipo Auxiliar correcto (o ninguno, en las líneas planas)."""
    client = flask_app.test_client()
    client.post("/login", data={"usuario": "", "clave": "test_local_only_1234"}, follow_redirects=True)

    r = client.post(
        "/conciliacion/procesar",
        data={"archivo": (io.BytesIO(_xlsx_cartola([
            (datetime(2026, 8, 18), "TRANSFERENCIA PAC HDI SEGUROS SA", 0, 103_733),
            (datetime(2026, 8, 20), "PAGO A PROVEEDOR ACME LTDA", 161_721, 0),
            (datetime(2026, 8, 21), "OF CENTRA", 50_000, 0),
        ])), "cartola.xlsx")},
        content_type="multipart/form-data",
    )
    check("procesar cartola -> 200", r.status_code == 200)

    form = {
        "archivo_nombre": "cartola.xlsx",
        "cuenta_banco_codigo": "1101-29",
        "cuenta_banco_descripcion": "BANCO BCI",
        "total_filas": "3",
        # Fila 0: abono, 1 línea (Clientes) con 1 documento (cubre el monto completo).
        "fecha_0": "18-08-2026", "fecha_iso_0": "2026-08-18", "detalle_0": "TRANSFERENCIA PAC HDI SEGUROS SA",
        "cargo_0": "0", "abono_0": "103733",
        "lineas_0_total": "1",
        "lineas_0_0_codigo": "1104-01", "lineas_0_0_descripcion": "DEUDORES CLIENTES", "lineas_0_0_monto": "",
        "lineas_0_0_docs_total": "1",
        "lineas_0_0_doc_0_modulo": "clientes", "lineas_0_0_doc_0_idx": "0",
        "lineas_0_0_doc_0_rut": "76123456-7", "lineas_0_0_doc_0_nombre": "PAC HDI SEGUROS SA",
        "lineas_0_0_doc_0_tipo_doc": "FAC-EL", "lineas_0_0_doc_0_numero_doc": "456",
        "lineas_0_0_doc_0_fecha_iso": "2026-08-18", "lineas_0_0_doc_0_monto": "103733",
        # Fila 1: cargo, 2 líneas — Proveedores con 1 documento (100.000) + cuenta suelta (61.721).
        "fecha_1": "20-08-2026", "fecha_iso_1": "2026-08-20", "detalle_1": "PAGO A PROVEEDOR ACME LTDA",
        "cargo_1": "161721", "abono_1": "0",
        "lineas_1_total": "2",
        "lineas_1_0_codigo": "2105-01", "lineas_1_0_descripcion": "FACTURAS POR PAGAR", "lineas_1_0_monto": "",
        "lineas_1_0_docs_total": "1",
        "lineas_1_0_doc_0_modulo": "proveedores", "lineas_1_0_doc_0_idx": "0",
        "lineas_1_0_doc_0_rut": "12345678-9", "lineas_1_0_doc_0_nombre": "ACME LTDA",
        "lineas_1_0_doc_0_tipo_doc": "FAC-EL", "lineas_1_0_doc_0_numero_doc": "789",
        "lineas_1_0_doc_0_fecha_iso": "2026-08-15", "lineas_1_0_doc_0_monto": "100000",
        "lineas_1_1_codigo": "5501-05", "lineas_1_1_descripcion": CUENTAS_POR_CODIGO["5501-05"]["descripcion"],
        "lineas_1_1_monto": "61721", "lineas_1_1_docs_total": "0",
        # Fila 2: cargo, 1 línea plana contra una cuenta suelta (sin documento).
        "fecha_2": "21-08-2026", "fecha_iso_2": "2026-08-21", "detalle_2": "OF CENTRA",
        "cargo_2": "50000", "abono_2": "0",
        "lineas_2_total": "1",
        "lineas_2_0_codigo": "5501-05", "lineas_2_0_descripcion": CUENTAS_POR_CODIGO["5501-05"]["descripcion"],
        "lineas_2_0_monto": "50000", "lineas_2_0_docs_total": "0",
        # Pool de documentos auxiliares ya cargados (lo que viajaría desde
        # los campos ocultos del fragmento de /cargar-auxiliar).
        "aux_clientes_total": "1",
        "aux_clientes_nombre_0": "PAC HDI SEGUROS SA", "aux_clientes_rut_0": "76123456-7",
        "aux_clientes_fecha_0": "01-08-2026", "aux_clientes_fecha_iso_0": "2026-08-01",
        "aux_clientes_tipo_documento_0": "FAC-EL", "aux_clientes_numero_documento_0": "456", "aux_clientes_monto_0": "103733",
        "aux_proveedores_total": "1",
        "aux_proveedores_nombre_0": "ACME LTDA", "aux_proveedores_rut_0": "12345678-9",
        "aux_proveedores_fecha_0": "15-08-2026", "aux_proveedores_fecha_iso_0": "2026-08-15",
        "aux_proveedores_tipo_documento_0": "FAC-EL", "aux_proveedores_numero_documento_0": "789", "aux_proveedores_monto_0": "100000",
        "aux_honorarios_total": "0",
    }

    r = client.post("/conciliacion/descargar", data=form)
    check("descargar -> 200, entrega un .xls", r.status_code == 200 and r.headers.get("Content-Type") == "application/vnd.ms-excel")

    wb = xlrd.open_workbook(file_contents=r.data)
    ws = wb.sheet_by_index(0)
    filas = [ws.row_values(i) for i in range(1, ws.nrows) if ws.cell(i, 4).value]  # columna Cuenta Detalle no vacía
    check("7 filas en total (2 + 3 + 2, banco incluido)", len(filas) == 7)

    total_debe = sum(f[8] for f in filas if isinstance(f[8], (int, float)))
    total_haber = sum(f[9] for f in filas if isinstance(f[9], (int, float)))
    check("el comprobante completo queda balanceado (Debe == Haber)", round(total_debe) == round(total_haber) == 103_733 + 161_721 + 50_000)

    fila_clientes = next(f for f in filas if f[4] == "1104-01")
    check("Clientes: Tipo Auxiliar 'A' en la línea 1104-01", fila_clientes[10] == "A")
    check("Clientes: Rut del documento", fila_clientes[11] == "76123456-7")
    check("Clientes: monto = 103.733 (documento completo)", fila_clientes[9] == 103_733.0)

    fila_proveedores = next(f for f in filas if f[4] == "2105-01")
    check("Proveedores: Tipo Auxiliar 'A' en la línea 2105-01", fila_proveedores[10] == "A")
    check("Proveedores: Rut del documento", fila_proveedores[11] == "12345678-9")
    check("Proveedores: monto = 100.000 (solo el documento, no el movimiento completo)", fila_proveedores[8] == 100_000.0)

    filas_gasto = [f for f in filas if f[4] == "5501-05"]
    check("2 líneas planas contra la cuenta de gasto (una por movimiento)", len(filas_gasto) == 2)
    check("ninguna línea plana lleva Tipo Auxiliar", all(not f[10] for f in filas_gasto))
    check("montos de las líneas planas: 61.721 (resto de fila 1) y 50.000 (fila 2)", {f[8] for f in filas_gasto} == {61_721.0, 50_000.0})

    # ---- Ahora un error de validación (falta la cuenta bancaria):
    # confirma que el pool de Clientes/Proveedores Y las líneas ya
    # armadas de cada movimiento se reconstruyen y se vuelven a mostrar
    # en pantalla, sin que el admin tenga que rehacer nada. ----
    form_incompleto = dict(form)
    form_incompleto["cuenta_banco_codigo"] = ""
    r = client.post("/conciliacion/descargar", data=form_incompleto)
    body = r.get_data(as_text=True)
    check("sin cuenta bancaria -> 200 (re-muestra la página con el error)", r.status_code == 200)
    check("el pool de Clientes se reconstruye en la página tras el error", "PAC HDI SEGUROS SA" in body)
    check("el pool de Proveedores se reconstruye en la página tras el error", "ACME LTDA" in body)
    check("el movimiento ya resuelto sigue mostrando su detalle", "TRANSFERENCIA PAC HDI SEGUROS SA" in body)
    check("las líneas ya armadas viajan de vuelta a la página (data-lineas-iniciales)", '&#34;codigo&#34;: &#34;1104-01&#34;' in body or '"codigo": "1104-01"' in body)


def main():
    test_procesar_fecha_guardada_como_texto()
    test_auxiliar_modulos_config()
    test_construir_filas_comprobantes_linea_plana()
    test_construir_filas_comprobantes_linea_con_varios_documentos()
    test_construir_filas_comprobantes_varias_lineas()
    test_cargar_auxiliar_rutas()
    test_descargar_con_lineas_multiples()

    print(f"\n{len(PASSED)} OK, {len(FAILED)} FAIL")
    if FAILED:
        print("Fallaron:", FAILED)
        sys.exit(1)


if __name__ == "__main__":
    main()
