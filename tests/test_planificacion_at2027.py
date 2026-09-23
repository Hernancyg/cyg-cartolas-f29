"""
Verificación de "Planificación AT 2027" (17-09-2026) — tabla editable de
seguimiento de empresas, más la carga masiva por Excel desde
Administrador. Mismo estilo sin-pytest que `tests/test_conciliacion.py`
(reusa el bootstrap de `tests/run_verification.py` — `flask_app`).

Uso:
    python tests/test_planificacion_at2027.py
"""

import io
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests.run_verification import flask_app, seed_data  # noqa: E402
from app.planificacion_at2027.pdf_calendario import generar_pdf_calendario  # noqa: E402

seed_data()  # crea el usuario "testuser"/"trabajador123" (rol trabajador) usado en test_trabajador_sin_acceso

PASSED, FAILED = [], []


def check(label, condition, extra=""):
    if condition:
        PASSED.append(label)
        print(f"  OK  {label}")
    else:
        FAILED.append(label)
        print(f" FAIL {label} {extra}")


def _login_admin(client):
    client.post("/login", data={"usuario": "", "clave": "test_local_only_1234"}, follow_redirects=True)


def _xlsx_planificacion(filas):
    """`filas`: lista de listas de 18 valores, mismo orden que
    `app/admin/routes.py:PLANIFICACION_COLUMNAS` (encabezado en la fila 1,
    datos desde la fila 2)."""
    wb = Workbook()
    ws = wb.active
    ws.append([
        "N°", "Empresa", "Analista", "Prioridad", "Caja/Banco",
        "Septiembre", "Octubre", "Noviembre", "Diciembre", "Enero", "Febrero",
        "Actualización Balance", "Reunión Cat1 (1°)", "Reunión Cat2", "Reunión Cat3",
        "Reunión Cat1 (2°)", "Grupo", "Estado Balance (Último mes trabajado)",
    ])
    for fila in filas:
        ws.append(fila)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_pagina_y_guardado_manual():
    client = flask_app.test_client()
    _login_admin(client)

    r = client.get("/planificacion_at2027/")
    check("GET /planificacion_at2027/ 200 (admin)", r.status_code == 200 and b"Planificaci" in r.data)

    r = client.get("/admin/planificacion_at2027")
    check("GET /admin/planificacion_at2027 200 (admin)", r.status_code == 200)
    check("subnav de Administrador incluye Planificación AT 2027", "Planificación AT 2027" in r.get_data(as_text=True))

    r = client.post("/planificacion_at2027/guardar", data={
        "p_numero": ["3", ""], "p_empresa": ["E DOS ASESORIA SPA", ""],
        "p_analista": ["David", ""], "p_prioridad": ["3", ""], "p_caja_banco": ["Caja", ""],
        "p_mes_septiembre": ["", ""], "p_mes_octubre": ["Javiera V", ""], "p_mes_noviembre": ["", ""],
        "p_mes_diciembre": ["Javiera V", ""], "p_mes_enero": ["Javiera V", ""], "p_mes_febrero": ["", ""],
        "p_actualizacion_balance": ["Semestral", ""], "p_reunion_cat1_1": ["", ""], "p_reunion_cat2": ["", ""],
        "p_reunion_cat3": ["", ""], "p_reunion_cat1_2": ["", ""], "p_grupo": ["", ""],
        "p_estado_balance_ultimo_mes": ["Diciembre", ""],
    }, follow_redirects=True)
    check("guardar manualmente -> 200", r.status_code == 200)
    body = r.get_data(as_text=True)
    check("la empresa guardada aparece de vuelta", "E DOS ASESORIA SPA" in body)
    check("la fila vacía (segunda) no se guarda", body.count("value=\"David\"") == 1)


def test_carga_masiva_valida():
    client = flask_app.test_client()
    _login_admin(client)

    xlsx = _xlsx_planificacion([
        [50, "ALIMENTOS SAN LUCAS", "Javier", 3, "Caja", "Catalina", "", "Catalina", "", "Javier", "", "Semestral", "", "", "", "", "", ""],
        [79, "INVERSIONES TBA SPA", "Javier", 1, "Banco", "", "Javier", "", "Javier", "", "Javier", "Mensual", "Octubre", "", "", "", "", "Febrero"],
    ])
    r = client.post(
        "/admin/planificacion_at2027/cargar",
        data={"archivo": (io.BytesIO(xlsx), "planificacion.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    check("carga masiva válida -> 200", r.status_code == 200)
    body = r.get_data(as_text=True)
    check("Planificación cargada: aviso de éxito", "cargada" in body)
    check("trae la primera empresa del Excel", "ALIMENTOS SAN LUCAS" in body)
    check("trae la segunda empresa del Excel", "INVERSIONES TBA SPA" in body)

    r = client.get("/planificacion_at2027/")
    body = r.get_data(as_text=True)
    check("la carga masiva REEMPLAZÓ lo guardado a mano (ya no aparece)", "E DOS ASESORIA SPA" not in body)


def test_carga_masiva_con_errores_no_guarda_nada():
    client = flask_app.test_client()
    _login_admin(client)

    xlsx = _xlsx_planificacion([
        [50, "ALIMENTOS SAN LUCAS", "Javier", 3, "Caja", "", "", "", "", "", "", "Semestral", "", "", "", "", "", ""],
        [None, "", "Silvana", 2, "Banco", "", "", "", "", "", "", "", "", "", "", "", "", ""],  # sin empresa
        [98, "SARAVIA Y GUZMAN SPA", "Jonathan", 9, "Banco", "", "", "", "", "", "", "", "", "", "", "", "", ""],  # prioridad inválida
    ])
    r = client.post(
        "/admin/planificacion_at2027/cargar",
        data={"archivo": (io.BytesIO(xlsx), "planificacion.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    check("carga con errores -> 200 (re-muestra la página, no revienta)", r.status_code == 200)
    body = r.get_data(as_text=True)
    check("avisa la fila sin 'Empresa'", "no tiene 'Empresa'" in body)
    check("avisa la fila con 'Prioridad' inválida", "no es 1, 2 o 3" in body)

    r = client.get("/planificacion_at2027/")
    body = r.get_data(as_text=True)
    check("nada se guardó: sigue la carga válida anterior, no la de esta prueba", "ALIMENTOS SAN LUCAS" in body and "SARAVIA Y GUZMAN SPA" not in body)


def test_estado_calculado_resumen_y_exportar():
    client = flask_app.test_client()
    _login_admin(client)

    # 4 filas cubriendo los 3 estados (19-09-2026: "Estado" se redefinió
    # de nuevo — ya no exige llegar a "Febrero" fijo, sino al ÚLTIMO mes
    # que de verdad tiene a alguien asignado en esa fila):
    #  - SIN AVANCE SPA: ningún mes asignado -> sin_asignar.
    #  - A MEDIO CAMINO SPA: asignada Sep/Oct/Nov (último = Noviembre),
    #    Avance = "Octubre" (no llega al último asignado) -> en_proceso.
    #  - TODO ASIGNADO SPA: asignada los 6 meses (último = Febrero),
    #    Avance = "Febrero" -> completado.
    #  - COMPLETADO ANTES DE FEBRERO SPA: asignada solo Sep/Oct (último =
    #    Octubre), Avance = "Octubre" -> completado, PESE A NO llegar a
    #    Febrero — esta fila es la que distingue la redefinición de la
    #    lógica anterior (que exigía Febrero sin importar hasta qué mes
    #    estaba realmente planificada la empresa).
    xlsx = _xlsx_planificacion([
        [1, "SIN AVANCE SPA", "David", 2, "Caja", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        [2, "A MEDIO CAMINO SPA", "Javier", 1, "Banco", "Javier", "Javier", "Javier", "", "", "", "Octubre", "", "", "", "", "", ""],
        [3, "TODO ASIGNADO SPA", "Catalina", 3, "Caja", "Catalina", "Catalina", "Catalina", "Catalina", "Catalina", "Catalina", "Febrero", "", "", "", "", "", ""],
        [4, "COMPLETADO ANTES DE FEBRERO SPA", "Marco", 2, "Banco", "Marco", "Marco", "", "", "", "", "Octubre", "", "", "", "", "", ""],
    ])
    client.post(
        "/admin/planificacion_at2027/cargar",
        data={"archivo": (io.BytesIO(xlsx), "planificacion.xlsx")},
        content_type="multipart/form-data",
    )

    r = client.get("/planificacion_at2027/")
    body = r.get_data(as_text=True)
    check("fila sin ningún mes asignado -> badge 'Sin asignar'", 'plan-estado-sin_asignar">' in body)
    check("fila cuyo avance no llega al último mes asignado -> 'En proceso'", 'plan-estado-en_proceso">' in body)
    check("fila con avance = Febrero (su último mes asignado) -> 'Completado'", 'plan-estado-completado">' in body)
    check(
        "fila con avance = Octubre pero SIN llegar a Febrero también -> 'Completado' "
        "(porque Octubre es su último mes asignado)",
        body.count('plan-estado-completado">') == 2,
    )
    check("tarjeta 'Total' = 4/4 (formato filtradas/total)", 'id="plan-kpi-total">4/4<' in body)
    check("tarjeta 'Completado' ya no se muestra (22-09-2026, pedido por el usuario: quitarla)", 'id="plan-kpi-completado"' not in body)
    check("tarjeta 'En proceso' cuenta 1 (25%)", 'id="plan-kpi-en-proceso">1 <span class="plan-kpi-pct">25%' in body)
    check("tarjeta 'Sin asignar' cuenta 1 (25%)", 'id="plan-kpi-sin-asignar">1 <span class="plan-kpi-pct">25%' in body)
    check("tarjeta 'Caja / Banco' presente", 'id="plan-kpi-caja"' in body and 'id="plan-kpi-banco"' in body)
    check("total por mes 'Sep' = 3 (A MEDIO CAMINO + TODO ASIGNADO + COMPLETADO ANTES)", 'id="plan-mes-sep-value">3<' in body)
    check("total por mes 'Ene' = 1 (solo TODO ASIGNADO)", 'id="plan-mes-ene-value">1<' in body)
    check("botón 'Mostrar reuniones' presente (columnas de Reunión ocultas por defecto)", 'id="plan-reuniones-btn"' in body)
    check("Avance Balance = Febrero queda seleccionado en su <select>", '<option value="Febrero" selected>Febrero</option>' in body)
    check("botón 'Generar informe PDF' presente", 'id="plan-informe-toggle-btn"' in body)
    check("botón 'Expandir' (18-09-2026: texto acortado)", 'id="plan-expandir-label">Expandir<' in body)
    check("filtro de Analista ya no menciona 'cualquier mes' (19-09-2026: solo mira esa columna)", "Analista (o cualquier mes)" not in body)
    check(
        "columna 'Estado Balance (último mes trabajado)' se acortó a 'Último cierre' (19-09-2026)",
        "Último cierre" in body and "Estado Balance (" not in body,
    )

    r = client.get("/planificacion_at2027/exportar")
    check("GET /planificacion_at2027/exportar -> 200", r.status_code == 200)
    check(
        "exportar entrega un .xlsx (Content-Type correcto)",
        "spreadsheetml.sheet" in (r.headers.get("Content-Type") or ""),
    )
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb.active
    encabezados = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    check("exportar: encabezado 'Empresa' en la columna 2", encabezados[1] == "Empresa")
    empresas_exportadas = [row[1].value for row in ws.iter_rows(min_row=2)]
    check("exportar: trae las 4 empresas cargadas", set(empresas_exportadas) == {
        "SIN AVANCE SPA", "A MEDIO CAMINO SPA", "TODO ASIGNADO SPA", "COMPLETADO ANTES DE FEBRERO SPA",
    })


def test_informe_pdf():
    client = flask_app.test_client()
    _login_admin(client)

    # Fila propia con "Avance Balance" y "Último cierre" bien distintos
    # (21-09-2026, pedido por el usuario: la columna "AVANCE" del informe
    # debe mostrar "estado_balance_ultimo_mes", no "actualizacion_balance")
    # — un upload propio en vez de depender de lo que dejó la prueba
    # anterior, para no acoplar ambas pruebas.
    xlsx = _xlsx_planificacion([
        [10, "EMPRESA JAVIER TEST SPA", "Javier", 1, "Caja", "Javier", "", "", "", "", "", "Septiembre", "", "", "", "", "", "Cierre 31-08-2026"],
        [11, "EMPRESA CATALINA TEST SPA", "Catalina", 2, "Banco", "", "Catalina", "", "", "", "", "Octubre", "", "", "", "", "", "Cierre 30-09-2026"],
    ])
    client.post(
        "/admin/planificacion_at2027/cargar",
        data={"archivo": (io.BytesIO(xlsx), "informe.xlsx")},
        content_type="multipart/form-data",
    )

    r = client.post("/planificacion_at2027/informe", data={}, follow_redirects=True)
    check("informe sin analistas elegidos -> 200 (re-muestra la página con el error)", r.status_code == 200)
    check("avisa que hay que elegir al menos un analista", "Selecciona al menos un analista" in r.get_data(as_text=True))

    r = client.post("/planificacion_at2027/informe", data={"analistas": ["Javier"]})
    check("informe de un analista -> 200", r.status_code == 200)
    check("informe entrega un .pdf (Content-Type correcto)", (r.headers.get("Content-Type") or "") == "application/pdf")
    check("informe: el PDF no viene vacío", len(r.data) > 500)

    with pdfplumber.open(io.BytesIO(r.data)) as pdf:
        texto_pdf = "\n".join(pagina.extract_text() or "" for pagina in pdf.pages)
    check("informe: la columna muestra el valor de 'Último cierre'", "Cierre 31-08-2026" in texto_pdf)
    check(
        "informe: NO muestra el valor de 'Avance Balance' (21-09-2026: se reemplazó por Último cierre)",
        "Septiembre" not in texto_pdf,
    )
    check("informe: el encabezado de esa columna dice 'ÚLTIMO CIERRE'", "ÚLTIMO CIERRE" in texto_pdf)

    r = client.post("/planificacion_at2027/informe", data={"analistas": ["Javier", "Catalina"]})
    check("informe de varios analistas a la vez -> 200", r.status_code == 200)
    check("informe entrega un .pdf (Content-Type correcto)", (r.headers.get("Content-Type") or "") == "application/pdf")


def test_calendario_reuniones_pdf():
    client = flask_app.test_client()
    _login_admin(client)

    # Una empresa con 2 categorías el mismo mes (Diciembre) para probar
    # que ambas etiquetas se muestran sin chocar, y otra con reuniones en
    # meses separados.
    xlsx = _xlsx_planificacion([
        [20, "EMPRESA CALENDARIO UNO SPA", "Pedro", 1, "Caja", "Pedro", "", "", "", "", "", "", "Diciembre", "Diciembre", "", "", "", ""],
        [21, "EMPRESA CALENDARIO DOS SPA", "Pedro", 2, "Banco", "", "Pedro", "", "", "", "", "", "", "", "Marzo-nunca", "Enero", "", ""],
    ])
    client.post(
        "/admin/planificacion_at2027/cargar",
        data={"archivo": (io.BytesIO(xlsx), "calendario.xlsx")},
        content_type="multipart/form-data",
    )

    r = client.post("/planificacion_at2027/informe/calendario", data={}, follow_redirects=True)
    check("calendario sin analistas elegidos -> 200 (re-muestra la página con el error)", r.status_code == 200)
    check("avisa que hay que elegir al menos un analista", "Selecciona al menos un analista" in r.get_data(as_text=True))

    r = client.post("/planificacion_at2027/informe/calendario", data={"analistas": ["Pedro"]})
    check("calendario de un analista -> 200", r.status_code == 200)
    check("calendario entrega un .pdf (Content-Type correcto)", (r.headers.get("Content-Type") or "") == "application/pdf")
    check("calendario: el PDF no viene vacío", len(r.data) > 500)

    with pdfplumber.open(io.BytesIO(r.data)) as pdf:
        texto_pdf = "\n".join(pagina.extract_text() or "" for pagina in pdf.pages)
    check("calendario: título del PDF presente", "CALENDARIO DE REUNIONES" in texto_pdf)
    check("calendario: nombre del analista presente", "Pedro" in texto_pdf)
    check("calendario: empresa con reunión en Diciembre aparece", "EMPRESA CALENDARIO UNO SPA" in texto_pdf)
    check("calendario: empresa con reunión en Enero aparece", "EMPRESA CALENDARIO DOS SPA" in texto_pdf)
    check("calendario: etiqueta de Cat.1 (1ª vez) presente", "1ª" in texto_pdf)
    check("calendario: etiqueta de Cat.2 presente", "2" in texto_pdf)
    check("calendario: etiqueta de Cat.1 (2ª vez) presente", "1ª·2" in texto_pdf)
    check(
        "calendario: un mes fuera del ciclo Sep-Feb (dato sucio) no revienta, solo se ignora",
        "Marzo-nunca" not in texto_pdf,
    )


def test_calendario_reuniones_pdf_pagina_larga():
    """Regresión (22-09-2026, bug real reportado por el usuario: "el
    informe sale cortado") — con muchas empresas juntas en un mismo mes,
    la lista de esa casilla podía ser más alta que una página COMPLETA;
    antes de la corrección, las últimas empresas se dibujaban igual y
    quedaban recortadas por el borde físico de la hoja en vez de pasar a
    una página nueva. Genera el PDF directo (sin pasar por Excel/HTTP,
    más simple para forzar ~20 empresas en un solo mes) y confirma que
    las 20 aparecen completas en el texto extraído, repartidas en más de
    una página."""
    nombres = [f"EMPRESA PRUEBA {i:02d} SPA" for i in range(1, 21)]
    filas = [{"empresa": n, "analista": "Constanza", "reunion_cat1_1": "Enero"} for n in nombres]

    pdf = generar_pdf_calendario([{"analista": "Constanza", "filas": filas}])
    data = pdf.getvalue()
    check("calendario largo: el PDF no viene vacío", len(data) > 500)

    with pdfplumber.open(io.BytesIO(data)) as doc:
        check("calendario largo: se reparte en más de 1 página", len(doc.pages) > 1)
        texto_pdf = "\n".join(pagina.extract_text() or "" for pagina in doc.pages)

    faltantes = [n for n in nombres if n not in texto_pdf]
    check("calendario largo: las 20 empresas aparecen completas (ninguna recortada)", not faltantes, faltantes)
    check("calendario largo: la columna continuada se marca '(CONT.)'", "(CONT.)" in texto_pdf)


def test_trabajador_sin_acceso():
    client = flask_app.test_client()
    client.post("/login", data={"usuario": "testuser", "clave": "trabajador123"}, follow_redirects=True)

    r = client.get("/planificacion_at2027/")
    check("trabajador NO puede ver Planificación AT 2027 (403)", r.status_code == 403)

    r = client.get("/planificacion_at2027/exportar")
    check("trabajador NO puede exportar a Excel (403)", r.status_code == 403)

    r = client.post("/planificacion_at2027/informe", data={"analistas": ["Javier"]})
    check("trabajador NO puede generar el informe PDF (403)", r.status_code == 403)

    r = client.post("/planificacion_at2027/informe/calendario", data={"analistas": ["Javier"]})
    check("trabajador NO puede generar el calendario de reuniones (403)", r.status_code == 403)

    r = client.get("/admin/planificacion_at2027")
    check("trabajador NO puede ver la carga masiva en Administrador (403)", r.status_code == 403)

    r = client.post("/admin/planificacion_at2027/cargar", data={}, content_type="multipart/form-data")
    check("trabajador NO puede cargar el Excel (403)", r.status_code == 403)


def main():
    test_pagina_y_guardado_manual()
    test_carga_masiva_valida()
    test_carga_masiva_con_errores_no_guarda_nada()
    test_estado_calculado_resumen_y_exportar()
    test_informe_pdf()
    test_calendario_reuniones_pdf()
    test_calendario_reuniones_pdf_pagina_larga()
    test_trabajador_sin_acceso()

    print(f"\n{len(PASSED)} OK, {len(FAILED)} FAIL")
    if FAILED:
        print("Fallaron:", FAILED)
        sys.exit(1)


if __name__ == "__main__":
    main()
