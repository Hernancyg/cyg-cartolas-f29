"""
Verificación de "Análisis" (22-09-2026) — genera el Excel de análisis
mensual a partir de los CSV que deja el puente Nubox -> repo en
`nubox_importado/`. Mismo estilo sin-pytest que `tests/test_depreciacion.py`
(reusa el bootstrap de `tests/run_verification.py` — `flask_app`).

Uso:
    python tests/test_analisis.py
"""

import csv
import io
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests.run_verification import flask_app, seed_data  # noqa: E402
from app.analisis import generador  # noqa: E402
from app.data import nubox_importado_repo  # noqa: E402

PASSED, FAILED = [], []


def check(label, condition, extra=""):
    if condition:
        PASSED.append(label)
        print(f"  OK  {label}")
    else:
        FAILED.append(label)
        print(f" FAIL {label} {extra}")


BALANCE_ENCABEZADOS = [
    "cuentaimputadaid", "cuenta_padre", "cuenta", "cuentaanalisis", "sucursal", "centrocosto",
    "debe", "haber", "deudor", "acreedor", "activo", "pasivo", "ingreso", "gasto", "tipocuentaid",
]
MAYOR_ENCABEZADOS = [
    "asientocontableid", "cuentaid", "cuenta", "codigocuenta", "descripcioncuenta", "fechamovimiento",
    "tipoasiento", "numero_asiento", "secuencia", "glosa", "sucursal", "centrocosto", "rutcontraparte",
    "contraparte", "descripciondetalle", "numerodocumento", "tipodocumento", "tipomovimiento",
    "fechadocumento", "debe", "haber", "saldofinal", "movimientocontableid",
]


def _escribir_csv(path, encabezados, filas):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(encabezados)
        for fila in filas:
            writer.writerow(fila)


def _fila_balance(cuenta, cuenta_padre, debe=0, haber=0, deudor=0, acreedor=0, activo=0, pasivo=0, ingreso=0, gasto=0):
    return ["1", cuenta_padre, cuenta, "", "", "", debe, haber, deudor, acreedor, activo, pasivo, ingreso, gasto, "1"]


def _fila_mayor(codigocuenta, descripcioncuenta, fechamovimiento, tipoasiento, numero_asiento, glosa, debe, haber, saldofinal):
    return [
        "", "", f"{codigocuenta} {descripcioncuenta}", codigocuenta, descripcioncuenta, fechamovimiento,
        tipoasiento, numero_asiento, "1", glosa, "", "", "", "", "", "", "", "", "",
        debe, haber, saldofinal, "",
    ]


def test_generar_excel_analisis(tmp_dir):
    ruta_balance = tmp_dir / "balance_general" / "TEST_202608.csv"
    ruta_mayor = tmp_dir / "libro_mayor" / "TEST_202608.csv"

    _escribir_csv(ruta_balance, BALANCE_ENCABEZADOS, [
        _fila_balance("1 ACTIVOS", ""),  # grupo, sin guión -> se descarta
        _fila_balance("1101-99 CAJA TEST", "1101 DISPONIBLE", debe=1_000_000, deudor=1_000_000, activo=1_000_000),
        _fila_balance("1101-88 BANCO CERO", "1101 DISPONIBLE"),  # hoja sin movimiento -> se descarta
        _fila_balance("2107-01 PROVEEDORES TEST", "2107 CTAS POR PAGAR", haber=400_000, acreedor=400_000, pasivo=400_000),
        _fila_balance("9999-01 CUENTA DESCONOCIDA", "9999 GRUPO RARO", debe=50_000, deudor=50_000, activo=50_000),
        _fila_balance("4101-01 VENTAS TEST", "4101 INGRESOS", haber=600_000, ingreso=600_000),
    ])

    _escribir_csv(ruta_mayor, MAYOR_ENCABEZADOS, [
        _fila_mayor("1101-99", "CAJA TEST", "", "", "", "Acumulado Anterior", 500_000, 0, 500_000),
        _fila_mayor("1101-99", "CAJA TEST", "2026-08-05T00:00:00.000Z", "Ingreso", "123", "Deposito", 500_000, 0, 1_000_000),
        _fila_mayor("2107-01", "PROVEEDORES TEST", "2026-08-10T00:00:00.000Z", "Egreso", "456", "Pago proveedor", 0, 400_000, -400_000),
    ])

    resultado = generador.generar_excel_analisis(ruta_balance, ruta_mayor, "EMPRESA DE PRUEBA SPA", "76.123.456-7", "202608")

    check("filas_mayor = 3 (Acumulado Anterior + 2 movimientos)", resultado.filas_mayor == 3, resultado.filas_mayor)
    check("cuentas_balance = 4 (se descartan la de grupo y la sin movimiento)", resultado.cuentas_balance == 4, resultado.cuentas_balance)
    check(
        "la cuenta de prefijo desconocido (9999) queda marcada sin clasificar",
        resultado.cuentas_sin_clasificar == ["9999-01 CUENTA DESCONOCIDA"],
        resultado.cuentas_sin_clasificar,
    )

    wb = openpyxl.load_workbook(io.BytesIO(resultado.archivo.getvalue()), keep_vba=True)

    ws_portada = wb["Portada"]
    check("Portada: nombre de la empresa", ws_portada["C6"].value == "EMPRESA DE PRUEBA SPA")
    check("Portada: RUT", ws_portada["C8"].value == "76.123.456-7")
    check("Portada: período legible", ws_portada["C14"].value == "AGOSTO 2026")

    ws_mayor = wb["Mayor"]
    check("Mayor: 3 filas de datos (filas 2-4), sin sobrantes de la plantilla", ws_mayor.max_row == 4)
    check("Mayor: fila 'Acumulado Anterior' toma el primer día del período como fecha", ws_mayor["C2"].value == "01-08-2026")
    check("Mayor: la fecha del segundo movimiento se convierte de ISO a dd-mm-aaaa", ws_mayor["C3"].value == "05-08-2026")
    check("Mayor: Código Cuenta en columna A", ws_mayor["A2"].value == "1101-99")
    check("Mayor: Cuenta (solo descripción, sin código) en columna B", ws_mayor["B2"].value == "CAJA TEST")
    check("Mayor: Saldo en columna L", ws_mayor["L4"].value == -400_000.0)

    ws_matriz = wb["Matriz"]
    check("Matriz: tantas filas de fórmulas como filas de Mayor (2-4)", ws_matriz.max_row == 4)
    check("Matriz: fórmula de la fila 4 referencia Mayor fila 4 (no quedó una fórmula vieja de otra fila)", ws_matriz["C4"].value == "=+Mayor!C4")

    ws_balance = wb["Balance"]
    check("Balance: primera cuenta hoja con movimiento en la fila 10", ws_balance["B10"].value == "1101-99 CAJA TEST")
    check("Balance: 4 cuentas -> filas 10-13, cierre en 14/15/16", ws_balance["B14"].value == "Sumas")
    # Sumas: Debe=1.000.000+50.000=1.050.000, Haber=400.000+600.000=1.000.000,
    # Deudor=1.050.000, Acreedor=400.000, Activo=1.050.000, Pasivo=400.000, Gasto=0, Ingreso=600.000.
    check("Balance: Sumas Debitos", ws_balance["C14"].value == 1_050_000)
    check("Balance: Sumas Creditos", ws_balance["D14"].value == 1_000_000)
    check("Balance: Sumas Ingresos", ws_balance["J14"].value == 600_000)
    check("Balance: Ganancia Ejercicio = Ingresos - Gastos = 600.000", ws_balance["H15"].value == 600_000 and ws_balance["I15"].value == 600_000)
    check("Balance: Totales Pasivo = Sumas Pasivo + Ganancia = 400.000+600.000", ws_balance["H16"].value == 1_000_000)
    check("Balance: Totales Perdidas = Sumas Perdidas + Ganancia = 0+600.000", ws_balance["I16"].value == 600_000)

    ws_clasif = wb["Clasificador"]
    # 4 cuentas en Balance, pero la 9999-01 queda sin clasificar -> Clasificador
    # trae solo 3 filas de fórmulas (filas 3-5), no 4.
    check("Clasificador: 3 filas (se salta la cuenta sin clasificar)", ws_clasif.max_row == 5, ws_clasif.max_row)
    check("Clasificador: offset +7 contra Balance (fila 3 -> Balance!B10)", ws_clasif["B3"].value == "=TRIM(Balance!B10)")
    check("Clasificador: categoría de la primera cuenta (1101-99 -> prefijo 1101)", ws_clasif["C3"].value == "Efectivo y equivalentes al efectivo")
    # Regresión del bug real: con offset fijo (+7 sobre la fila de Clasificador)
    # esta fila terminaba apuntando a Balance!B12 en vez de B13, porque
    # 9999-01 (la 3ra cuenta de Balance, fila 12) se salta por no tener
    # clasificación conocida. 4101-01 es la 4ta cuenta de Balance (fila
    # 10+3=13) pero solo la 3ra fila de Clasificador (fila 5).
    check("Clasificador: cuenta después de un salto referencia la fila de Balance correcta (no corrida)", ws_clasif["B5"].value == "=TRIM(Balance!B13)", ws_clasif["B5"].value)
    check("Clasificador: categoría de la cuenta después del salto (4101-01 -> prefijo 4101)", ws_clasif["C5"].value == "Resultado del Ejercicio", ws_clasif["C5"].value)


def test_nubox_importado_repo_lee_datos_reales():
    """Usa los CSV reales que ya dejó el puente Nubox -> repo (empresa 626,
    período 202608) para confirmar que el repo de listado los encuentra."""
    disponibles = nubox_importado_repo.listar_disponibles()
    check(
        "el repo encuentra la combinación 626/202608 que ya trajo el puente",
        {"alias": "626", "periodo": "202608"} in disponibles,
        disponibles,
    )
    rutas = nubox_importado_repo.rutas_de("626", "202608")
    check("rutas_de encuentra ambos archivos reales", rutas is not None and all(p.is_file() for p in rutas))
    check("rutas_de(alias inexistente) devuelve None", nubox_importado_repo.rutas_de("no-existe", "999999") is None)


def test_rutas_flujo_completo():
    client = flask_app.test_client()
    client.post("/login", data={"usuario": "", "clave": "test_local_only_1234"}, follow_redirects=True)

    r = client.get("/analisis/")
    check("GET /analisis/ -> 200", r.status_code == 200)
    check("la página lista la empresa 626 real ya disponible", "626" in r.get_data(as_text=True))

    r = client.post(
        "/analisis/generar",
        data={"alias": "626", "periodo": "202608", "empresa_nombre": "SOCIEDAD AGRICOLA Y COMERCIAL ACEVEDO Y COMPANIA SPA", "empresa_rut": "76.353.060-4"},
    )
    check("generar con datos reales -> 200", r.status_code == 200)
    check(
        "descarga un .xlsm (Content-Type correcto)",
        (r.headers.get("Content-Type") or "") == "application/vnd.ms-excel.sheet.macroEnabled.12",
    )
    check("el archivo generado no viene vacío", len(r.data) > 10_000)

    r = client.post("/analisis/generar", data={"alias": "no-existe", "periodo": "999999", "empresa_nombre": "X"}, follow_redirects=True)
    check("combinación inexistente -> avisa el error, no revienta", r.status_code == 200 and "no se encontraron" in r.get_data(as_text=True).lower())

    r = client.post("/analisis/generar", data={"alias": "626", "periodo": "202608", "empresa_nombre": ""}, follow_redirects=True)
    check("sin nombre de empresa -> avisa el error", r.status_code == 200 and "nombre de la empresa" in r.get_data(as_text=True).lower())


def test_trabajador_no_puede_ver_analisis():
    client = flask_app.test_client()
    r = client.post("/login", data={"usuario": "testuser", "clave": "trabajador123"}, follow_redirects=True)
    check("login trabajador ok (fixture)", r.status_code == 200)
    r = client.get("/analisis/")
    check("trabajador NO puede ver Análisis por defecto (403)", r.status_code == 403)
    r = client.post("/analisis/generar", data={"alias": "626", "periodo": "202608", "empresa_nombre": "X"})
    check("trabajador NO puede generar el Excel (403)", r.status_code == 403)


def main():
    import tempfile

    seed_data()  # crea el usuario "testuser"/"trabajador123" (rol trabajador) usado en test_trabajador_no_puede_ver_analisis

    with tempfile.TemporaryDirectory() as tmp:
        test_generar_excel_analisis(Path(tmp))

    test_nubox_importado_repo_lee_datos_reales()
    test_rutas_flujo_completo()
    test_trabajador_no_puede_ver_analisis()

    print(f"\n{len(PASSED)} OK, {len(FAILED)} FAIL")
    if FAILED:
        print("Fallaron:", FAILED)
        sys.exit(1)


if __name__ == "__main__":
    main()
