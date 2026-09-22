"""
Genera el Excel de análisis mensual (Estado de Situación Financiera
Clasificado) a partir de una plantilla `.xlsm` con macros/fórmulas
(`app/static/plantillas/Analisis_plantilla.xlsm`) y los CSV de Balance
General / Libro Mayor que deja el puente Nubox -> repo en
`nubox_importado/` (ver `app/data/nubox_importado_repo.py`).

Cómo está armada la plantilla (ver también `app/analisis/clasificacion_
cuentas.json` y el plan en `.claude/plans/`): de las 11 hojas, solo dos
son datos crudos que hay que pegar (`Mayor`, `Balance`); todo lo demás
(`Matriz`, `Clasificador`, `EEFF`, y las 5 hojas de nota) son fórmulas
que se recalculan solas al abrir el archivo en Excel. `Matriz` tiene una
fila de fórmulas por CADA fila de `Mayor` (indexada por número de fila,
literal) y `Clasificador` tiene una fila de fórmulas por CADA cuenta
CLASIFICADA de `Balance` — por eso este módulo regenera esas dos hojas
fila por fila en vez de solo pegar valores. Ojo con `Clasificador`: como
salta las cuentas sin clasificación conocida, su número de fila NO
queda alineado 1:1 con el de `Balance` (un offset fijo se rompe apenas
se salta la primera cuenta) — hay que calcular la fila de `Balance` de
cada cuenta por su POSICIÓN real en `cuentas_balance`, no por la fila
que le tocó en `Clasificador` (ver `_regenerar_clasificador`, bug real
encontrado y corregido el 22-09-2026).

La columna "Cuenta EEFF" de `Clasificador` (a qué categoría IFRS
pertenece cada cuenta) NO es derivable matemáticamente del código de
cuenta — es un criterio contable que el usuario ya aplicó una vez sobre
un cliente real (TLINK SPA). Se extrajo esa clasificación a
`clasificacion_cuentas.json` (código exacto de cuenta -> categoría, más
un respaldo por los primeros 4 dígitos del código para cuentas nuevas
que compartan el mismo prefijo) — funciona bien mientras las empresas
usen el mismo plan de cuentas de la firma. Una cuenta que no calce con
ninguna de las dos tablas queda fuera de `Clasificador` y se reporta
como advertencia (`GeneracionResultado.cuentas_sin_clasificar`).
"""

import copy
import csv
import json
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import List, Optional

import openpyxl

PLANTILLA_PATH = Path(__file__).resolve().parent.parent / "static" / "plantillas" / "Analisis_plantilla.xlsm"
CLASIFICACION_PATH = Path(__file__).resolve().parent / "clasificacion_cuentas.json"

MESES_LABEL = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]

_BALANCE_PRIMERA_FILA_DATOS = 10
_CLASIFICADOR_PRIMERA_FILA_DATOS = 3


@dataclass
class GeneracionResultado:
    archivo: BytesIO
    filas_mayor: int
    cuentas_balance: int
    cuentas_sin_clasificar: List[str] = field(default_factory=list)


def _leer_csv(path: Path) -> List[dict]:
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def _to_float(valor) -> float:
    if valor is None or valor == "":
        return 0.0
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0.0


def _cargar_clasificacion() -> dict:
    with open(CLASIFICACION_PATH, encoding="utf-8") as fh:
        return json.load(fh)


def _clasificar_cuenta(codigo: str, mapa: dict) -> Optional[str]:
    if codigo in mapa["por_codigo"]:
        return mapa["por_codigo"][codigo]
    prefijo = codigo.split("-")[0]
    return mapa["por_prefijo"].get(prefijo)


def _fecha_mayor(fecha_iso: str, fecha_periodo_inicio: str) -> str:
    """Convierte 'aaaa-mm-ddT00:00:00.000Z' (Nubox) a texto 'dd-mm-aaaa'
    (mismo formato — texto, no fecha real de Excel — que usa la
    plantilla). Las filas 'Acumulado Anterior' de Nubox no traen fecha;
    se les pone el primer día del período para que `Matriz` pueda
    calcular YEAR()/MONTH() sin reventar."""
    if not fecha_iso:
        return fecha_periodo_inicio
    try:
        return datetime.strptime(fecha_iso[:10], "%Y-%m-%d").strftime("%d-%m-%Y")
    except ValueError:
        return fecha_periodo_inicio


def _periodo_a_fecha_inicio(periodo_yyyymm: str) -> str:
    anio, mes = periodo_yyyymm[:4], periodo_yyyymm[4:6]
    return f"01-{mes}-{anio}"


def periodo_legible(periodo_yyyymm: str) -> str:
    anio, mes = int(periodo_yyyymm[:4]), int(periodo_yyyymm[4:6])
    return f"{MESES_LABEL[mes - 1]} {anio}"


def _limpiar_copiar_estilo(ws, fila_origen: int, fila_destino: int, num_columnas: int):
    """Copia el estilo (formato de celda) de `fila_origen` a `fila_destino`
    — para que filas nuevas agregadas más allá de las que trae la
    plantilla se vean igual que las de ejemplo, no con el formato por
    defecto de Excel."""
    for c in range(1, num_columnas + 1):
        origen = ws.cell(fila_origen, c)
        destino = ws.cell(fila_destino, c)
        if origen.has_style:
            destino._style = copy.copy(origen._style)


def _escribir_mayor(ws, filas_csv: List[dict], fecha_inicio: str) -> int:
    # Limpia cualquier fila de ejemplo que traiga la plantilla, dejando
    # solo el encabezado (fila 1).
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    fila = 2
    for f in filas_csv:
        ws.cell(fila, 1, f.get("codigocuenta") or "")
        ws.cell(fila, 2, f.get("descripcioncuenta") or "")
        ws.cell(fila, 3, _fecha_mayor(f.get("fechamovimiento") or "", fecha_inicio))
        ws.cell(fila, 4, f.get("tipoasiento") or "")
        ws.cell(fila, 5, f.get("numero_asiento") or "")
        ws.cell(fila, 6, f.get("secuencia") or "")
        ws.cell(fila, 7, f.get("glosa") or "")
        ws.cell(fila, 8, f.get("centrocosto") or "")
        ws.cell(fila, 9, f.get("sucursal") or "")
        debe = _to_float(f.get("debe"))
        haber = _to_float(f.get("haber"))
        ws.cell(fila, 10, debe if debe else None)
        ws.cell(fila, 11, haber if haber else None)
        ws.cell(fila, 12, _to_float(f.get("saldofinal")))
        fila += 1
    return fila - 2  # cantidad de filas de datos escritas


def _formula_matriz(r: int) -> list:
    return [
        f"=+YEAR(C{r})",
        f"=+MONTH(C{r})",
        f"=+Mayor!C{r}",
        f"=+Mayor!E{r}",
        f"=+Mayor!A{r}",
        f"=+Mayor!J{r}",
        f"=+Mayor!K{r}",
        f"=+Mayor!G{r}",
        f'=IF(D{r}="00000000",0,+F{r}-G{r})',
        f'=IF(+MID(E{r},1,1)="1","Activo",IF(+MID(E{r},1,1)="2","Pasivo",IF(+MID(E{r},1,1)="3","Pasivo","EERR")))',
        f'=+IF(J{r}="Activo",F{r}-G{r},-F{r}+G{r})',
        f'=+Mayor!A{r}&" "&Mayor!B{r}',
    ]


def _regenerar_matriz(ws, filas_mayor: int):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in range(2, filas_mayor + 2):
        for col, formula in enumerate(_formula_matriz(r), start=1):
            ws.cell(r, col, formula)


def _escribir_balance(ws, filas_csv: List[dict]) -> List[dict]:
    """Filtra las cuentas HOJA (código con guión) que tienen movimiento
    real (algún valor ≠ 0 en los 8 campos), las escribe desde la fila 10,
    y calcula las 3 filas de cierre (Sumas/Ganancia Ejercicio/Totales)
    exactamente como lo hacía la planilla real del usuario (ver el
    módulo docstring). Devuelve la lista de cuentas escritas (para poder
    clasificarlas después en `Clasificador`)."""
    CAMPOS = ("debe", "haber", "deudor", "acreedor", "activo", "pasivo", "gasto", "ingreso")

    def es_hoja_con_movimiento(f):
        codigo = (f.get("cuenta") or "").split(" ", 1)[0]
        if "-" not in codigo:
            return False
        return any(abs(_to_float(f.get(c))) > 0.0001 for c in CAMPOS)

    cuentas = [f for f in filas_csv if es_hoja_con_movimiento(f)]

    fila = _BALANCE_PRIMERA_FILA_DATOS
    for f in cuentas:
        ws.cell(fila, 2, f.get("cuenta") or "")
        ws.cell(fila, 3, _to_float(f.get("debe")))
        ws.cell(fila, 4, _to_float(f.get("haber")))
        ws.cell(fila, 5, _to_float(f.get("deudor")))
        ws.cell(fila, 6, _to_float(f.get("acreedor")))
        ws.cell(fila, 7, _to_float(f.get("activo")))
        ws.cell(fila, 8, _to_float(f.get("pasivo")))
        ws.cell(fila, 9, _to_float(f.get("gasto")))
        ws.cell(fila, 10, _to_float(f.get("ingreso")))
        fila += 1

    sumas = {col: 0.0 for col in range(3, 11)}
    for r in range(_BALANCE_PRIMERA_FILA_DATOS, fila):
        for col in range(3, 11):
            sumas[col] += ws.cell(r, col).value or 0.0

    fila_sumas = fila
    fila_ganancia = fila + 1
    fila_totales = fila + 2

    ws.cell(fila_sumas, 2, "Sumas")
    for col in range(3, 11):
        ws.cell(fila_sumas, col, round(sumas[col], 2))

    ganancia_ejercicio = round(sumas[10] - sumas[9], 2)  # Ganancias(J) - Perdidas(I)
    ws.cell(fila_ganancia, 2, "Ganancia Ejercicio")
    ws.cell(fila_ganancia, 8, ganancia_ejercicio)  # Pasivo
    ws.cell(fila_ganancia, 9, ganancia_ejercicio)  # Perdidas

    ws.cell(fila_totales, 2, "Totales")
    ws.cell(fila_totales, 3, round(sumas[3], 2))
    ws.cell(fila_totales, 4, round(sumas[4], 2))
    ws.cell(fila_totales, 5, round(sumas[5], 2))
    ws.cell(fila_totales, 6, round(sumas[6], 2))
    ws.cell(fila_totales, 7, round(sumas[7], 2))
    ws.cell(fila_totales, 8, round(sumas[8] + ganancia_ejercicio, 2))
    ws.cell(fila_totales, 9, round(sumas[9] + ganancia_ejercicio, 2))
    ws.cell(fila_totales, 10, round(sumas[10], 2))

    # La plantilla trae "Sumas"/"Ganancia Ejercicio"/"Totales" en filas
    # fijas (137-139, pensadas para las 127 cuentas de ejemplo) — si esta
    # empresa tiene menos cuentas, esas filas quedan más abajo que el
    # bloque de cierre recién escrito y hay que borrarlas para no dejar un
    # segundo "Sumas"/"Totales" vacío y confuso más abajo en la hoja.
    if ws.max_row > fila_totales:
        ws.delete_rows(fila_totales + 1, ws.max_row - fila_totales)

    return cuentas


def _regenerar_clasificador(ws, cuentas_balance: List[dict], mapa_clasificacion: dict) -> List[str]:
    """Una fila de fórmulas por cada cuenta CLASIFICADA de `Balance`.

    OJO (bug real encontrado en verificación, 22-09-2026): la fila de
    `Balance` que le corresponde a cada cuenta es SIEMPRE
    `_BALANCE_PRIMERA_FILA_DATOS + i` (`i` = posición de esa cuenta
    dentro de `cuentas_balance`, el mismo orden en que `_escribir_balance`
    las escribió) — NO se puede derivar sumando un offset fijo al número
    de fila de `Clasificador`, porque `Clasificador` salta las cuentas
    sin clasificación conocida y por lo tanto sus filas dejan de estar
    alineadas 1:1 con las de `Balance` apenas se salta la primera cuenta
    sin clasificar. Con el offset fijo, cada cuenta clasificada DESPUÉS
    de un salto terminaba apuntando a la cuenta de `Balance` equivocada
    (mismo número de filas, pero corridas), sumando plata de una cuenta a
    la categoría de otra sin ningún error visible en Excel. Las cuentas
    sin clasificación conocida se devuelven para avisarle al usuario."""
    if ws.max_row >= _CLASIFICADOR_PRIMERA_FILA_DATOS:
        ws.delete_rows(_CLASIFICADOR_PRIMERA_FILA_DATOS, ws.max_row - _CLASIFICADOR_PRIMERA_FILA_DATOS + 1)

    sin_clasificar = []
    fila_clasificador = _CLASIFICADOR_PRIMERA_FILA_DATOS
    for i, f in enumerate(cuentas_balance):
        codigo = (f.get("cuenta") or "").split(" ", 1)[0]
        categoria = _clasificar_cuenta(codigo, mapa_clasificacion)
        if categoria is None:
            sin_clasificar.append(f.get("cuenta") or codigo)
            continue
        fila_balance = _BALANCE_PRIMERA_FILA_DATOS + i
        ws.cell(fila_clasificador, 2, f"=TRIM(Balance!B{fila_balance})")
        ws.cell(fila_clasificador, 3, categoria)
        ws.cell(fila_clasificador, 4, f"=+SUMIF(Matriz!E:E,MID(Clasificador!B{fila_clasificador},1,7),Matriz!I:I)")
        fila_clasificador += 1
    return sin_clasificar


def _escribir_portada(ws, empresa_nombre: str, empresa_rut: str, periodo_yyyymm: str):
    ws["C6"] = empresa_nombre
    ws["C8"] = empresa_rut
    ws["C14"] = periodo_legible(periodo_yyyymm)


def generar_excel_analisis(
    ruta_balance: Path, ruta_mayor: Path, empresa_nombre: str, empresa_rut: str, periodo_yyyymm: str,
) -> GeneracionResultado:
    filas_balance_csv = _leer_csv(ruta_balance)
    filas_mayor_csv = _leer_csv(ruta_mayor)
    mapa_clasificacion = _cargar_clasificacion()
    fecha_inicio_periodo = _periodo_a_fecha_inicio(periodo_yyyymm)

    wb = openpyxl.load_workbook(PLANTILLA_PATH, keep_vba=True)

    _escribir_portada(wb["Portada"], empresa_nombre, empresa_rut, periodo_yyyymm)
    filas_mayor = _escribir_mayor(wb["Mayor"], filas_mayor_csv, fecha_inicio_periodo)
    _regenerar_matriz(wb["Matriz"], filas_mayor)
    cuentas_balance = _escribir_balance(wb["Balance"], filas_balance_csv)
    sin_clasificar = _regenerar_clasificador(wb["Clasificador"], cuentas_balance, mapa_clasificacion)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return GeneracionResultado(
        archivo=buffer, filas_mayor=filas_mayor, cuentas_balance=len(cuentas_balance),
        cuentas_sin_clasificar=sin_clasificar,
    )
