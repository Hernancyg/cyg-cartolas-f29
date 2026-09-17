"""
"Administrador": puerto de `pagina_administrador()` de la app de
Streamlit. Ahora edita directamente las tablas de Supabase (`cuentas_config`
/ `usuarios`) en vez de descargar/subir JSON a GitHub — la mejora natural
que permite tener una base de datos real detrás.
"""

import io

from flask import Blueprint, render_template, request, redirect, url_for, flash, current_app

from openpyxl import load_workbook

from app.auth.decorators import admin_required, login_required
from app.auth.security import hash_password
from app.parsers.config_manager import cargar_config, guardar_config, CuentaConfig
from app.parsers.f29_parser import parsear_f29, _clean_monto
from app.data import usuarios_repo, visibilidad_repo, tipos_documento_repo, planificacion_at2027_repo
from app.nav import paginas_con_visibilidad

admin_bp = Blueprint("admin", __name__, url_prefix="/admin")

SENTINEL_SIN_OPERADOR = "(sin símbolo)"
OPERADORES = ["+", "-", "=", SENTINEL_SIN_OPERADOR]


# ---------------------------------------------------------------------------
# Cuentas F29
# ---------------------------------------------------------------------------

def _leer_filas_config(form):
    cuentas = form.getlist("c_cuenta")
    codigos = form.getlist("c_codigo")
    descripciones = form.getlist("c_descripcion")
    tipos = form.getlist("c_tipo")
    anclas = form.getlist("c_ancla")
    operadores = form.getlist("c_operador")

    configs = []
    for i in range(len(cuentas)):
        campos = [cuentas[i], codigos[i], descripciones[i], tipos[i], anclas[i]]
        if any(c is None or str(c).strip() == "" for c in campos):
            continue
        operador_val = operadores[i].strip() if i < len(operadores) else ""
        if operador_val == SENTINEL_SIN_OPERADOR:
            operador_val = ""
        configs.append(CuentaConfig(
            cuenta=cuentas[i].strip(), codigo_f29=codigos[i].strip(),
            descripcion=descripciones[i].strip(), tipo=tipos[i].strip(),
            texto_ancla=anclas[i].strip(), operador=operador_val,
        ))
    return configs


@admin_bp.route("/cuentas", methods=["GET"])
@admin_required
def cuentas():
    configs = cargar_config()
    return render_template(
        "admin/cuentas.html", configs=configs, sentinel=SENTINEL_SIN_OPERADOR,
        operadores=OPERADORES, prueba=None,
    )


@admin_bp.route("/cuentas/guardar", methods=["POST"])
@admin_required
def cuentas_guardar():
    nuevas_configs = _leer_filas_config(request.form)
    guardar_config(nuevas_configs)
    flash("Cambios guardados. Se aplicarán de inmediato.", "success")
    return redirect(url_for("admin.cuentas"))


@admin_bp.route("/cuentas/probar", methods=["POST"])
@admin_required
def cuentas_probar():
    nuevas_configs = _leer_filas_config(request.form)
    archivo = request.files.get("pdf_prueba")
    prueba = None
    if archivo and archivo.filename and nuevas_configs:
        try:
            data_prueba = parsear_f29(io.BytesIO(archivo.read()), configs=nuevas_configs)
            prueba = {
                "filas": [{
                    "cuenta": c.cuenta, "codigo_f29": c.codigo_f29,
                    "monto": _clean_monto(data_prueba.valores.get(c.codigo_f29)),
                    "detectado": data_prueba.valores.get(c.codigo_f29) is not None,
                } for c in nuevas_configs],
                "advertencias": data_prueba.advertencias,
            }
        except Exception as exc:  # noqa: BLE001
            flash(f"No se pudo leer el PDF de prueba: {exc}", "error")
    return render_template(
        "admin/cuentas.html", configs=nuevas_configs or cargar_config(),
        sentinel=SENTINEL_SIN_OPERADOR, operadores=OPERADORES, prueba=prueba,
    )


# ---------------------------------------------------------------------------
# Usuarios
# ---------------------------------------------------------------------------

@admin_bp.route("/usuarios", methods=["GET"])
@admin_required
def usuarios():
    from flask import session
    lista = usuarios_repo.listar_usuarios()
    yo = session.get("usuario") or {}
    return render_template(
        "admin/usuarios.html", usuarios=lista, roles=usuarios_repo.ROLES,
        mi_id=yo.get("id"),
    )


@admin_bp.route("/usuarios/actualizar", methods=["POST"])
@admin_required
def usuarios_actualizar():
    from flask import session
    ids = request.form.getlist("u_id")
    nombres = request.form.getlist("u_nombre")
    roles = request.form.getlist("u_rol")
    activos = set(request.form.getlist("u_activo"))  # ids marcados activos
    mi_id = (session.get("usuario") or {}).get("id")

    for i, uid in enumerate(ids):
        if uid == mi_id:
            continue  # un admin no puede editarse su propio rol/estado, para no auto-bloquearse
        usuarios_repo.actualizar_datos(
            uid, nombres[i] if i < len(nombres) else "",
            roles[i] if i < len(roles) else "trabajador",
            uid in activos,
        )
    flash("Cambios de usuarios guardados.", "success")
    return redirect(url_for("admin.usuarios"))


@admin_bp.route("/usuarios/crear", methods=["POST"])
@admin_required
def usuarios_crear():
    usuario = (request.form.get("usuario") or "").strip()
    nombre = (request.form.get("nombre") or "").strip()
    rol = request.form.get("rol") or "trabajador"
    clave = request.form.get("clave") or ""
    clave_confirmar = request.form.get("clave_confirmar") or ""

    if not usuario or not nombre or not clave:
        flash("Completa usuario, nombre y contraseña.", "error")
    elif clave != clave_confirmar:
        flash("Las contraseñas no coinciden.", "error")
    elif len(clave) < 4:
        flash("La contraseña debe tener al menos 4 caracteres.", "error")
    elif usuarios_repo.buscar_usuario(usuario) is not None:
        flash(f"Ya existe un usuario '{usuario}'.", "error")
    else:
        usuarios_repo.crear_usuario(usuario, nombre, rol, clave)
        flash(f"Usuario '{usuario}' creado.", "success")
    return redirect(url_for("admin.usuarios"))


@admin_bp.route("/usuarios/resetear", methods=["POST"])
@admin_required
def usuarios_resetear():
    usuario_id = request.form.get("usuario_id")
    nueva_clave = request.form.get("nueva_clave") or ""
    if len(nueva_clave) < 4:
        flash("La contraseña debe tener al menos 4 caracteres.", "error")
    elif usuario_id:
        usuarios_repo.restablecer_password(usuario_id, nueva_clave)
        flash("Contraseña actualizada.", "success")
    return redirect(url_for("admin.usuarios"))


@admin_bp.route("/usuarios/eliminar", methods=["POST"])
@admin_required
def usuarios_eliminar():
    from flask import session
    usuario_id = request.form.get("usuario_id")
    mi_id = (session.get("usuario") or {}).get("id")
    if usuario_id and usuario_id != mi_id:
        usuarios_repo.eliminar_usuario(usuario_id)
        flash("Usuario eliminado.", "success")
    else:
        flash("No puedes eliminar tu propio usuario.", "error")
    return redirect(url_for("admin.usuarios"))


# ---------------------------------------------------------------------------
# Visibilidad de pestañas (09-09-2026): activar cada pestaña del menú para
# todos los usuarios o dejarla solo para administradores, sin tocar código
# ni redesplegar. "Administrador" no aparece aquí a propósito — es la
# pantalla que administra esto mismo, así que se queda siempre solo-admin.
# ---------------------------------------------------------------------------

@admin_bp.route("/pestanas", methods=["GET"])
@admin_required
def pestanas():
    configurables = [p for p in paginas_con_visibilidad() if p["configurable"]]
    return render_template("admin/pestanas.html", paginas=configurables)


@admin_bp.route("/pestanas/guardar", methods=["POST"])
@admin_required
def pestanas_guardar():
    configurables = [p for p in paginas_con_visibilidad() if p["configurable"]]
    mapa = {}
    for p in configurables:
        valor = request.form.get(f"vis_{p['endpoint']}", "admin")
        mapa[p["endpoint"]] = (valor == "admin")
    try:
        visibilidad_repo.guardar_todas(mapa)
    except Exception as exc:  # noqa: BLE001
        # La causa más común es que todavía no se haya ejecutado
        # `migration/003_visibilidad_pestanas.sql` en Supabase (la tabla
        # `visibilidad_pestanas` no existe). En vez de un Internal Server
        # Error sin explicación, se avisa qué falta y no se pierde nada:
        # la lectura (`obtener_mapa`) ya tolera esto y sigue usando los
        # valores por defecto de `app/nav.py`.
        current_app.logger.warning("No se pudo guardar visibilidad_pestanas: %s", exc)
        flash(
            "No se pudo guardar: falta crear la tabla 'visibilidad_pestanas' en Supabase "
            "(ejecuta migration/003_visibilidad_pestanas.sql una vez en el SQL Editor) o "
            "hubo un problema de conexión. Mientras tanto, las pestañas siguen con su "
            "configuración de siempre.",
            "error",
        )
        return redirect(url_for("admin.pestanas"))
    flash("Visibilidad de pestañas guardada. Se aplica de inmediato.", "success")
    return redirect(url_for("admin.pestanas"))


# ---------------------------------------------------------------------------
# Tipos de Documento (10-09-2026): mapeo de texto ("FAC-EL", "BOL-HE", etc.)
# a código numérico, usado por "Empresas Caja" para la columna "Tipo De
# Documento" del bloque de Tipo Auxiliar "A"/"H" del archivo de salida —
# ver `app/data/tipos_documento_repo.py`.
# ---------------------------------------------------------------------------

@admin_bp.route("/tipos_documento", methods=["GET"])
@admin_required
def tipos_documento():
    mapa = tipos_documento_repo.obtener_mapa()
    filas = sorted(mapa.items())
    return render_template("admin/tipos_documento.html", filas=filas)


@admin_bp.route("/tipos_documento/guardar", methods=["POST"])
@admin_required
def tipos_documento_guardar():
    textos = request.form.getlist("td_texto")
    codigos = request.form.getlist("td_codigo")
    mapa = {}
    for i in range(len(textos)):
        texto = (textos[i] or "").strip()
        codigo = (codigos[i] if i < len(codigos) else "").strip()
        if not texto or not codigo:
            continue
        try:
            mapa[texto] = int(codigo)
        except ValueError:
            flash(f"El código de “{texto}” debe ser un número entero — no se guardó esa fila.", "error")
    try:
        tipos_documento_repo.guardar_todos(mapa)
    except Exception as exc:  # noqa: BLE001
        current_app.logger.warning("No se pudo guardar tipos_documento: %s", exc)
        flash(
            "No se pudo guardar: falta crear la tabla 'tipos_documento' en Supabase "
            "(ejecuta migration/004_tipos_documento.sql una vez en el SQL Editor) o "
            "hubo un problema de conexión. Mientras tanto, siguen los códigos de siempre.",
            "error",
        )
        return redirect(url_for("admin.tipos_documento"))
    flash("Tipos de Documento guardados. Se aplican de inmediato.", "success")
    return redirect(url_for("admin.tipos_documento"))


# ---------------------------------------------------------------------------
# Planificación AT 2027 (17-09-2026): carga masiva por Excel que reemplaza
# de una vez toda la tabla que edita la pestaña "Planificación AT 2027"
# (ver `app/planificacion_at2027/routes.py` y `app/data/
# planificacion_at2027_repo.py`) — mismo criterio "el Excel reemplaza
# todo" confirmado con el usuario, sin intentar preservar ediciones
# manuales frente a una carga nueva.
# ---------------------------------------------------------------------------

PLANIFICACION_ALLOWED_EXT = (".xlsx", ".xlsm")

# Orden de columnas del Excel de carga masiva — igual orden que la tabla
# en pantalla (encabezado en la fila 1, datos desde la fila 2).
PLANIFICACION_COLUMNAS = [
    "numero", "empresa", "analista", "prioridad", "caja_banco",
    "mes_septiembre", "mes_octubre", "mes_noviembre", "mes_diciembre", "mes_enero", "mes_febrero",
    "actualizacion_balance", "reunion_cat1_1", "reunion_cat2", "reunion_cat3", "reunion_cat1_2",
    "grupo", "estado_balance_ultimo_mes",
]


def _texto_celda(valor):
    if valor is None:
        return ""
    return str(valor).strip()


def _leer_excel_planificacion(file_storage):
    """Devuelve (filas, errores). Si `errores` no está vacía, `filas` es
    `None` — no se guarda nada hasta que el Excel se corrija (mismo
    criterio "todo o nada" que `app/conciliacion/routes.py:descargar`,
    agrupando los números de fila reales del Excel con el mismo problema
    en un solo mensaje en vez de uno por fila)."""
    try:
        wb = load_workbook(io.BytesIO(file_storage.read()), data_only=True)
    except Exception as exc:  # noqa: BLE001
        return None, [f"No se pudo abrir el archivo: {exc}"]

    ws = wb.active
    total_columnas = len(PLANIFICACION_COLUMNAS)
    filas = []
    filas_sin_empresa = []
    filas_prioridad_invalida = []

    for numero_fila, fila_excel in enumerate(
        ws.iter_rows(min_row=2, max_col=total_columnas, values_only=True), start=2,
    ):
        valores = list(fila_excel) + [None] * (total_columnas - len(fila_excel))
        if all(_texto_celda(v) == "" for v in valores):
            continue  # fila vacía (ej. al final de la hoja) — se ignora sin avisar

        empresa = _texto_celda(valores[1])
        if not empresa:
            filas_sin_empresa.append(numero_fila)
            continue

        prioridad_texto = _texto_celda(valores[3])
        prioridad = None
        if prioridad_texto:
            if prioridad_texto in ("1", "2", "3"):
                prioridad = int(prioridad_texto)
            else:
                filas_prioridad_invalida.append(numero_fila)

        numero_texto = _texto_celda(valores[0])
        numero = int(numero_texto) if numero_texto.isdigit() else None

        fila = {
            "numero": numero, "empresa": empresa, "analista": _texto_celda(valores[2]),
            "prioridad": prioridad, "caja_banco": _texto_celda(valores[4]),
        }
        for i, columna in enumerate(PLANIFICACION_COLUMNAS[5:], start=5):
            fila[columna] = _texto_celda(valores[i])
        filas.append(fila)

    errores = []
    if filas_sin_empresa:
        etiqueta = "fila" if len(filas_sin_empresa) == 1 else "filas"
        errores.append(
            f"La {etiqueta} {', '.join(str(n) for n in filas_sin_empresa)} del Excel no tiene "
            "'Empresa' — corrígela y vuelve a subir el archivo."
        )
    if filas_prioridad_invalida:
        etiqueta = "fila" if len(filas_prioridad_invalida) == 1 else "filas"
        errores.append(
            f"La {etiqueta} {', '.join(str(n) for n in filas_prioridad_invalida)} tiene una "
            "'Prioridad' que no es 1, 2 o 3 — corrígela y vuelve a subir el archivo."
        )

    if errores:
        return None, errores
    return filas, []


@admin_bp.route("/planificacion_at2027", methods=["GET"])
@admin_required
def planificacion_at2027():
    return render_template("admin/planificacion_at2027.html")


@admin_bp.route("/planificacion_at2027/cargar", methods=["POST"])
@admin_required
def planificacion_at2027_cargar():
    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        flash("Sube un archivo Excel para continuar.", "error")
        return redirect(url_for("admin.planificacion_at2027"))

    if not archivo.filename.lower().endswith(PLANIFICACION_ALLOWED_EXT):
        flash("Formato no permitido. Sube un archivo .xlsx o .xlsm.", "error")
        return redirect(url_for("admin.planificacion_at2027"))

    filas, errores = _leer_excel_planificacion(archivo)
    if errores:
        for mensaje in errores:
            flash(mensaje, "error")
        return redirect(url_for("admin.planificacion_at2027"))
    if not filas:
        flash("El archivo no trae ninguna fila con datos.", "error")
        return redirect(url_for("admin.planificacion_at2027"))

    try:
        planificacion_at2027_repo.guardar_todos(filas)
    except Exception as exc:  # noqa: BLE001
        current_app.logger.warning("No se pudo guardar planificacion_at2027 (carga masiva): %s", exc)
        flash(
            "No se pudo guardar: falta crear la tabla 'planificacion_at2027' en Supabase "
            "(ejecuta migration/014_planificacion_at2027.sql una vez en el SQL Editor) o "
            "hubo un problema de conexión.",
            "error",
        )
        return redirect(url_for("admin.planificacion_at2027"))

    flash(f"Planificación cargada: {len(filas)} empresas.", "success")
    return redirect(url_for("planificacion_at2027.index"))
