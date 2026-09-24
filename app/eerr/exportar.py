"""
Exportación del "EERR Dinámico" a Excel (openpyxl) y PDF (reportlab puro,
mismo criterio que `app/planificacion_at2027/pdf_informe.py`: sin
dependencias de sistema para Render). Ambos reciben el informe ya calculado
por `app.eerr.calculo.calcular`, así que muestran las mismas cifras que la
pantalla, con el estilo de grilla elegido por el usuario (opción C, al
estilo de la vista de Nubox): encabezado gris claro, conceptos en banda
suave, cuentas en azul, totales y márgenes en banda gris azulada,
negativos entre paréntesis y celdas vacías cuando no hay monto.
"""

from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle

BRAND_ICON = Path(__file__).resolve().parent.parent / "static" / "img" / "brand-icon.png"

# Paleta de la grilla (opción C)
HEAD = "F3F5F8"
LINE = "E7EAF0"
CONCEPTO = "F7F8FA"
TOTAL = "E2E8F0"
MARGEN = "D8E1EC"
CUENTA = "1F3F73"
TINTA = "1B2433"
GRIS = "8A94A6"
NEG = "A3303A"
AVISO = "FDF1DC"
AVISO_TXT = "8A5A00"

_FILL = {"grupo": CONCEPTO, "manual": CONCEPTO, "total": TOTAL, "formula": MARGEN, "resultado": MARGEN, "sinasig": AVISO}


def _titulo_columna_acum(informe, anio):
    cols = informe["columnas"]
    if len(cols) == 1:
        return f"{cols[0][:3]} {anio}"
    return f"{cols[0][:3]} a {cols[-1][:3]} {anio}"


def _filas_visibles(informe, con_cuentas):
    return [f for f in informe["filas"] if con_cuentas or f["tipo"] != "cta"]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def generar_excel(informe: dict, empresa: dict, anio: int, rango: str, fuente: str, con_cuentas: bool = True) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Estado de Resultados"
    ws.sheet_view.showGridLines = False

    n = len(informe["columnas"])
    ultima_col = 1 + 2 * (n + 1)
    fmt_monto = '#,##0;(#,##0);""'
    fmt_pct = '0.0%;(0.0%);""'
    borde = Border(left=Side(style="thin", color=LINE), right=Side(style="thin", color=LINE),
                   top=Side(style="thin", color=LINE), bottom=Side(style="thin", color=LINE))

    ws["A1"] = "Estado de Resultados"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=TINTA)
    ws["A2"] = empresa.get("razon_social") or ""
    ws["A2"].font = Font(name="Calibri", size=11, bold=True, color=TINTA)
    detalle = [f"RUT {empresa['rut']}" if empresa.get("rut") else "", rango, f"Fuente: {fuente}"]
    ws["A3"] = " · ".join(d for d in detalle if d)
    ws["A3"].font = Font(name="Calibri", size=9, color=GRIS)

    fila_enc = 5
    encabezados = ["CONCEPTO"]
    for c in informe["columnas"]:
        encabezados += [f"{c.upper()} {anio}", "%"]
    encabezados += [_titulo_columna_acum(informe, anio).upper(), "%"]
    for j, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_enc, column=j, value=texto)
        celda.font = Font(name="Calibri", size=8, bold=True, color=TINTA if texto != "%" else GRIS)
        celda.fill = PatternFill("solid", fgColor=HEAD)
        celda.alignment = Alignment(horizontal="left" if j == 1 else "right", vertical="center")
        celda.border = borde

    r = fila_enc + 1
    for f in _filas_visibles(informe, con_cuentas):
        tipo = f["tipo"]
        es_cta = tipo == "cta"
        color = CUENTA if es_cta else (AVISO_TXT if tipo == "sinasig" else TINTA)
        negrita = not es_cta
        tam = 9 if es_cta else 8
        fill = PatternFill("solid", fgColor=_FILL[tipo]) if tipo in _FILL else None
        nombre = ("    " + f["nombre"]) if es_cta else (f["nombre"].upper() if tipo == "resultado" else f["nombre"])
        valores = []
        for v, p in zip(f["valores"], f["pcts"]):
            valores += [(v, fmt_monto), (None if p is None else p / 100, fmt_pct)]
        valores += [(f["acum"], fmt_monto), (None if f["acum_pct"] is None else f["acum_pct"] / 100, fmt_pct)]

        c0 = ws.cell(row=r, column=1, value=nombre)
        c0.font = Font(name="Calibri", size=tam, bold=negrita, color=color)
        c0.alignment = Alignment(vertical="center", wrap_text=True)
        c0.border = borde
        if fill:
            c0.fill = fill
        for j, (valor, fmt) in enumerate(valores, start=2):
            es_pct = fmt == fmt_pct
            celda = ws.cell(row=r, column=j, value=round(valor) if (valor and not es_pct) else (valor or None))
            celda.number_format = fmt
            negativo = isinstance(valor, (int, float)) and valor < 0
            celda.font = Font(name="Calibri", size=8 if es_pct else tam, bold=negrita and not es_pct,
                              color=NEG if negativo else (GRIS if es_pct else color))
            celda.alignment = Alignment(horizontal="right", vertical="center")
            celda.border = borde
            if fill:
                celda.fill = fill
        r += 1

    ws.column_dimensions["A"].width = 44
    for j in range(2, ultima_col + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14 if j % 2 == 0 else 7.5
    ws.freeze_panes = ws.cell(row=fila_enc + 1, column=2)
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila_enc}:{fila_enc}"

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _monto_pdf(v):
    if not v:
        return ""
    texto = "{:,.0f}".format(abs(v)).replace(",", ".")
    return f"({texto})" if v < 0 else texto


def _pct_pdf(p):
    if p is None:
        return ""
    texto = "{:.1f}%".format(abs(p)).replace(".", ",")
    return f"({texto})" if p < 0 else texto


def generar_pdf(informe: dict, empresa: dict, anio: int, rango: str, fuente: str, con_cuentas: bool = True) -> BytesIO:
    n = len(informe["columnas"])
    pagina = landscape(A4) if n <= 6 else landscape(A3)
    margen = 12 * mm
    ancho = pagina[0] - 2 * margen
    ancho_concepto = 62 * mm
    ancho_pct = 10 * mm
    ancho_monto = (ancho - ancho_concepto - (n + 1) * ancho_pct) / (n + 1)

    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=pagina, leftMargin=margen, rightMargin=margen,
                            topMargin=margen, bottomMargin=margen,
                            title=f"Estado de Resultados · {empresa.get('razon_social', '')}")
    hex_ = lambda h: colors.HexColor("#" + h)  # noqa: E731
    est_titulo = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=14, textColor=hex_(TINTA), leading=17)
    est_sub = ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=10, textColor=hex_(TINTA), leading=13)
    est_meta = ParagraphStyle("m", fontName="Helvetica", fontSize=8, textColor=hex_(GRIS), leading=10)
    est_nombre = ParagraphStyle("n", fontName="Helvetica", fontSize=6.5, leading=8)

    historia = []
    cabecera = [
        Paragraph("Estado de Resultados", est_titulo),
        Paragraph(empresa.get("razon_social") or "", est_sub),
        Paragraph(" · ".join(d for d in [f"RUT {empresa['rut']}" if empresa.get("rut") else "", rango,
                                         f"Fuente: {fuente}", f"Emitido el {datetime.now():%d-%m-%Y}"] if d), est_meta),
    ]
    if BRAND_ICON.exists():
        # El ícono es blanco con fondo transparente: va sobre un cuadro azul
        # marino, igual que en los PDF de Planificación AT 2027.
        logo = Image(str(BRAND_ICON), width=9 * mm, height=9 * mm)
        marco = Table([[logo]], colWidths=[12 * mm], rowHeights=[12 * mm], style=[
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0A1224")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ])
        historia.append(Table([[marco, cabecera]], colWidths=[16 * mm, ancho - 16 * mm],
                              style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    else:
        historia.extend(cabecera)
    historia.append(Spacer(1, 5 * mm))

    encabezado = ["CONCEPTO"]
    for c in informe["columnas"]:
        encabezado += [f"{c.upper()} {anio}", "%"]
    encabezado += [_titulo_columna_acum(informe, anio).upper(), "%"]
    datos = [encabezado]
    estilos = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 6),
        ("TEXTCOLOR", (0, 0), (-1, 0), hex_(TINTA)),
        ("BACKGROUND", (0, 0), (-1, 0), hex_(HEAD)),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, hex_(LINE)),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    for k in range(2, len(encabezado), 2):
        estilos.append(("TEXTCOLOR", (k, 0), (k, 0), hex_(GRIS)))

    for f in _filas_visibles(informe, con_cuentas):
        i = len(datos)
        tipo = f["tipo"]
        es_cta = tipo == "cta"
        color = CUENTA if es_cta else (AVISO_TXT if tipo == "sinasig" else TINTA)
        fuente_fila = "Helvetica" if es_cta else "Helvetica-Bold"
        tam = 6.5 if es_cta else 6
        nombre = f["nombre"].upper() if tipo == "resultado" else f["nombre"]
        est = ParagraphStyle(f"r{i}", parent=est_nombre, fontName=fuente_fila, fontSize=tam,
                             textColor=hex_(color), leftIndent=8 if es_cta else 0)
        fila = [Paragraph(nombre.replace("&", "&amp;").replace("<", "&lt;"), est)]
        for v, p in zip(f["valores"], f["pcts"]):
            fila += [_monto_pdf(v), _pct_pdf(p)]
        fila += [_monto_pdf(f["acum"]), _pct_pdf(f["acum_pct"])]
        datos.append(fila)

        estilos.append(("FONT", (1, i), (-1, i), fuente_fila, tam))
        estilos.append(("TEXTCOLOR", (1, i), (-1, i), hex_(color)))
        for k in range(2, len(fila), 2):
            estilos.append(("FONT", (k, i), (k, i), "Helvetica", 5.5))
            estilos.append(("TEXTCOLOR", (k, i), (k, i), hex_(GRIS)))
        montos = f["valores"] + [f["acum"]]
        for m, v in enumerate(montos):
            if v and v < 0:
                estilos.append(("TEXTCOLOR", (1 + 2 * m, i), (2 + 2 * m, i), hex_(NEG)))
        if tipo in _FILL:
            estilos.append(("BACKGROUND", (0, i), (-1, i), hex_(_FILL[tipo])))

    anchos = [ancho_concepto] + [ancho_monto, ancho_pct] * (n + 1)
    tabla = Table(datos, colWidths=anchos, repeatRows=1)
    tabla.setStyle(TableStyle(estilos))
    historia.append(tabla)

    def pie(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(hex_(GRIS))
        canvas.drawString(margen, 6 * mm, "C&G Group · Estado de Resultados Dinámico")
        canvas.drawRightString(pagina[0] - margen, 6 * mm, f"Página {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(historia, onFirstPage=pie, onLaterPages=pie)
    salida.seek(0)
    return salida
