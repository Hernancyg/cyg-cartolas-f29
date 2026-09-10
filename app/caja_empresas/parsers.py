"""
Parsers de los 3 archivos de entrada de "Empresas Caja" (Clientes,
Proveedores, Honorarios) — estados de cuenta que el usuario exporta desde
su sistema contable, con documentos pendientes de cobro o pago.

Mapeo de columnas confirmado por el usuario contra 3 archivos de ejemplo
reales (`Clientes_2.xlsx`, `Proveedores_2.xlsx`, `Honorarios_2.xlsx`):

- Clientes / Proveedores: mismo formato ("Estado de Cuentas - Pendientes"),
  encabezado en la fila 2, datos desde la fila 3. Columna E = Nombre,
  F = Fecha Registro (datetime), K = Tipo Documento, L = N° Documento,
  P = Debe (monto a cobrar, Clientes), Q = Haber (monto a pagar,
  Proveedores).
- Honorarios ("Estado de Cuentas de Honorario - Pendientes"): encabezado en
  la fila 3, datos desde la fila 4. Columna B = Nombre, C = Fecha (texto
  DD-MM-AAAA, no datetime), F = Boleta (por ejemplo "BOL-HE 2" — se separa
  en tipo de documento "BOL-HE" + número "2"), J = Haber (monto). Las filas
  de subtotal ("Total Prestador") y el total final ("Total Informe") no
  tienen Rut (columna A) y se descartan por eso.

Cada fila devuelta ya viene con `seleccionado: True` por defecto — el
admin destilda en pantalla los documentos que no van en esta cobranza/pago,
en vez de tener que marcar uno por uno.
"""

from datetime import datetime

from openpyxl import load_workbook


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor.strip()
    return str(valor).strip()


def _monto(valor):
    try:
        if valor in (None, ""):
            return 0.0
        return abs(float(valor))
    except (TypeError, ValueError):
        return 0.0


def _abrir(file_storage):
    import io
    try:
        wb = load_workbook(io.BytesIO(file_storage.read()), data_only=True)
        return wb.active, None
    except Exception as exc:  # noqa: BLE001
        return None, f"No se pudo abrir el archivo: {exc}"


def _parsear_clientes_proveedores(file_storage, col_monto):
    """`col_monto`: 16 (P, Clientes/Debe) o 17 (Q, Proveedores/Haber)."""
    ws, error = _abrir(file_storage)
    if error:
        return None, error

    filas = []
    for fila in ws.iter_rows(min_row=3, values_only=False):
        nombre = _texto(fila[4].value if len(fila) > 4 else None)  # col E
        if not nombre:
            continue
        fecha = fila[5].value if len(fila) > 5 else None  # col F
        tipo_documento = _texto(fila[10].value if len(fila) > 10 else None)  # col K
        numero_documento = _texto(fila[11].value if len(fila) > 11 else None)  # col L
        monto = _monto(fila[col_monto - 1].value if len(fila) >= col_monto else None)

        if not isinstance(fecha, datetime) or monto <= 0:
            continue

        filas.append({
            "nombre": nombre,
            "fecha": fecha.strftime("%d-%m-%Y"),
            "fecha_iso": fecha.strftime("%Y-%m-%d"),
            "tipo_documento": tipo_documento,
            "numero_documento": numero_documento,
            "monto": monto,
            "cuenta_codigo": "",
            "cuenta_descripcion": "",
            "seleccionado": True,
        })
    return filas, None


def parsear_clientes(file_storage):
    """Documentos por cobrar (columna P = Debe). Devuelve (filas, error)."""
    return _parsear_clientes_proveedores(file_storage, col_monto=16)


def parsear_proveedores(file_storage):
    """Documentos por pagar (columna Q = Haber). Devuelve (filas, error)."""
    return _parsear_clientes_proveedores(file_storage, col_monto=17)


def parsear_honorarios(file_storage):
    """Boletas de honorarios pendientes de pago. Devuelve (filas, error)."""
    ws, error = _abrir(file_storage)
    if error:
        return None, error

    filas = []
    for fila in ws.iter_rows(min_row=4, values_only=False):
        rut = _texto(fila[0].value if len(fila) > 0 else None)  # col A
        if not rut:
            continue  # subtotal "Total Prestador" / total "Total Informe" / fila vacía
        nombre = _texto(fila[1].value if len(fila) > 1 else None)  # col B
        fecha_txt = _texto(fila[2].value if len(fila) > 2 else None)  # col C
        boleta = _texto(fila[5].value if len(fila) > 5 else None)  # col F
        monto = _monto(fila[9].value if len(fila) > 9 else None)  # col J

        try:
            fecha = datetime.strptime(fecha_txt, "%d-%m-%Y")
        except ValueError:
            continue  # fecha no reconocible, no se puede armar el comprobante

        if monto <= 0:
            continue

        if " " in boleta:
            tipo_documento, numero_documento = boleta.rsplit(" ", 1)
        else:
            tipo_documento, numero_documento = boleta, ""

        filas.append({
            "nombre": nombre,
            "fecha": fecha.strftime("%d-%m-%Y"),
            "fecha_iso": fecha.strftime("%Y-%m-%d"),
            "tipo_documento": tipo_documento,
            "numero_documento": numero_documento,
            "monto": monto,
            "cuenta_codigo": "",
            "cuenta_descripcion": "",
            "seleccionado": True,
        })
    return filas, None
