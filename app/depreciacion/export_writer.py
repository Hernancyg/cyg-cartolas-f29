"""
Genera el Excel de la tabla de depreciación de una empresa para un
período dado (una fila por activo, ver `app/depreciacion/calculo.py`).
Mismo estilo simple (Arial, encabezado en negrita, bordes finos) que
`app/parsers/output_writer.py`.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "Activo", "Fecha Adquisición", "Valor Adquisición", "Vida Útil (años)",
    "Depreciación Mensual", "Meses Depreciados", "Depreciación Acumulada",
    "Valor Libro",
]

COL_WIDTHS = {"A": 40, "B": 16, "C": 18, "D": 14, "E": 18, "F": 16, "G": 20, "H": 14}
NUMFMT_MONTO = "#,##0"
NUMFMT_FECHA = "dd-mm-yyyy"

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_tabla_workbook(empresa_nombre: str, periodo_label: str, filas: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Depreciación"

    ws["A1"] = f"{empresa_nombre} — Depreciación {periodo_label}"
    ws["A1"].font = Font(name="Arial", size=10, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    header_row = 3
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    for i, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=i, value=header)
        cell.font = Font(name="Arial", size=8, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    row = header_row + 1
    for fila in filas:
        valores = [
            fila["nombre_activo"] + (" (dado de baja)" if fila.get("de_baja") else ""),
            fila["fecha_adquisicion"],
            fila["valor_adquisicion"],
            fila["vida_util_anios"],
            fila["depreciacion_mensual"],
            f"{fila['meses_transcurridos']} / {fila['meses_vida_util']}",
            fila["depreciacion_acumulada"],
            fila["valor_libro"],
        ]
        for col_idx, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row, column=col_idx, value=valor)
            cell.font = Font(name="Arial", size=8)
            cell.border = BORDER
            col_letter = get_column_letter(col_idx)
            if col_letter == "B":
                cell.number_format = NUMFMT_FECHA
                cell.alignment = Alignment(horizontal="center")
            elif col_letter in ("C", "E", "G", "H"):
                cell.number_format = NUMFMT_MONTO
                cell.alignment = Alignment(horizontal="right")
            elif col_letter in ("D", "F"):
                cell.alignment = Alignment(horizontal="center")
        row += 1

    if filas:
        total_row = row
        ws.cell(row=total_row, column=1, value="Total").font = Font(name="Arial", size=8, bold=True)
        for col_letter, key in (("C", "valor_adquisicion"), ("E", "depreciacion_mensual"), ("G", "depreciacion_acumulada"), ("H", "valor_libro")):
            total = sum(f[key] for f in filas)
            cell = ws[f"{col_letter}{total_row}"]
            cell.value = total
            cell.number_format = NUMFMT_MONTO
            cell.font = Font(name="Arial", size=8, bold=True)
            cell.alignment = Alignment(horizontal="right")
            cell.border = BORDER

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


KARDEX_HEADERS = [
    "Fecha", "Costo Inicial", "Adiciones", "Costo Total", "Deprec. Acum. (apertura)",
    "Valor Libro", "Vida Útil (meses)", "Vida Útil Consumida (meses)",
    "Depreciación del Ejercicio", "Depreciación Acumulada Total",
]
KARDEX_COL_WIDTHS = {"A": 14, "B": 16, "C": 14, "D": 16, "E": 20, "F": 14, "G": 16, "H": 20, "I": 20, "J": 22}


def build_kardex_workbook(empresa_nombre: str, activo: dict, filas: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kardex"

    ws["A1"] = f"{empresa_nombre} — {activo['nombre_activo']} — Kardex de depreciación"
    ws["A1"].font = Font(name="Arial", size=10, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(KARDEX_HEADERS))

    header_row = 3
    for col, width in KARDEX_COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    for i, header in enumerate(KARDEX_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=i, value=header)
        cell.font = Font(name="Arial", size=8, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    row = header_row + 1
    for fila in filas:
        valores = [
            fila["fecha"], activo["valor_adquisicion"], fila["adiciones"], fila["costo_total"],
            fila["deprec_acum_apertura"], fila["valor_libro"], fila["vida_util_antes_meses"],
            fila["meses_utilizados"], fila["depreciacion_ejercicio"], fila["deprec_acum_cierre"],
        ]
        for col_idx, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row, column=col_idx, value=valor)
            cell.font = Font(name="Arial", size=8)
            cell.border = BORDER
            col_letter = get_column_letter(col_idx)
            if col_letter == "A":
                cell.number_format = NUMFMT_FECHA
                cell.alignment = Alignment(horizontal="center")
            elif col_letter in ("B", "C", "D", "E", "F", "I", "J"):
                cell.number_format = NUMFMT_MONTO
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="center")
        row += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
